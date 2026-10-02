from __future__ import annotations

import asyncio
import importlib.util
import json
import shutil
import subprocess
import sys
import time
from contextlib import asynccontextmanager
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import AsyncMock, Mock

import pytest
from aiohttp import web
from aiohttp.test_utils import make_mocked_request

from bot.plugins.artifacts import ArtifactStore
from bot.plugins.host_api import PluginHostApi
from bot.plugins.manifest import load_plugin_manifest
from bot.plugins.registry import PluginRegistry
from bot.plugins.runtime import PluginRuntime
from bot.web.api_common import AuthContext, WebApiError
from bot.web.auth_store import CAP_RUN_PLUGINS
from bot.web.server import WebApiServer

PLUGIN = Path(__file__).resolve().parents[1] / "examples/plugins/html-docx"
HTML = b'<!doctype html><html><head><title>Report</title></head><body><h1>Original</h1><p>Text</p></body></html>'
STUB_CORE = '''
def inspect_html(source, resources=None, config=None):
    return {"title":"Report", "paragraphCount":2, "tableCount":0, "imageCount":0,
            "warnings":[], "errors":[], "resources": []}
def export_docx(source, resources=None, config=None):
    return b"worker-export", inspect_html(source, resources, config)
'''


@pytest.fixture
def converter_available():
    assert (PLUGIN / "backend/converter.py").is_file()
    for name in ("docx", "lxml", "tinycss2", "cssselect2", "PIL", "resvg_py"):
        pytest.importorskip(name)


@asynccontextmanager
async def copied_plugin(tmp_path, *, config=None, limits=None, core=None, missing=False, host_type=PluginHostApi):
    plugin_dir = tmp_path / "installed plugins/html-docx"
    shutil.copytree(PLUGIN, plugin_dir, ignore=shutil.ignore_patterns("__pycache__"))
    if core is not None:
        (plugin_dir / "backend/converter.py").write_text(core, encoding="utf-8")
    if core is not None or missing:
        worker = plugin_dir / "backend/worker.py"
        override = "importlib.util.find_spec = lambda _: None" if missing else "DEPENDENCIES = {}"
        worker.write_text(worker.read_text(encoding="utf-8").replace(
            'if __name__ == "__main__":', override + '\n\nif __name__ == "__main__":'), encoding="utf-8")
    manifest_path = plugin_dir / "plugin.json"
    raw = json.loads(manifest_path.read_text(encoding="utf-8"))
    raw["config"].update(config or {})
    if limits is not None:
        raw["runtime"]["limits"] = limits
    manifest_path.write_text(json.dumps(raw), encoding="utf-8")
    workspace = tmp_path / "workspace"
    workspace.mkdir()
    source = workspace / "sample.html"
    source.write_bytes(HTML)
    artifacts = ArtifactStore(tmp_path)
    audit = []
    runtime = PluginRuntime(workspace_root_for=lambda _: workspace,
                            host_api=host_type(artifacts), audit_hook=audit.append)
    try:
        yield SimpleNamespace(root=plugin_dir, workspace=workspace, source=source, artifacts=artifacts,
                              audit=audit, runtime=runtime, manifest=load_plugin_manifest(manifest_path))
    finally:
        await runtime.shutdown()


async def open_html(plugin, path="sample.html"):
    return await plugin.runtime.open_view("main", plugin.manifest, "html-docx", {"path": path})


async def export_html(plugin, view):
    return await plugin.runtime.invoke_action("main", plugin.manifest, view_id="html-docx",
                                             session_id=view["sessionId"], action_id="export-docx", payload={})


def png(color):
    from PIL import Image
    output = BytesIO()
    Image.new("RGB", (32, 16), color).save(output, "PNG")
    return output.getvalue()


@pytest.mark.asyncio
async def test_copy_install_rereads_source_and_images_and_downloads_docx(tmp_path, converter_available):
    from docx import Document
    async with copied_plugin(tmp_path) as plugin:
        registry = PluginRegistry(plugin.root.parent)
        for name in ("sample.HTML", "sample.htm"):
            resolution = registry.resolve_file_handler(name)
            assert (resolution.plugin_id, resolution.view_id) == ("html-docx", "html-docx")
        assert plugin.manifest.runtime.permissions.workspace_read
        assert plugin.manifest.runtime.permissions.temp_artifacts
        assert not plugin.manifest.runtime.permissions.workspace_list
        plugin.source.write_bytes(HTML.replace(b"</body>", b'<img src="image.png"></body>'))
        image = plugin.workspace / "image.png"
        image.write_bytes(png("red"))
        opened = await open_html(plugin)
        assert not opened["summary"]["actions"][0]["disabled"]
        assert opened["initialWindow"]["rows"][0]["cells"]["imageCount"] == 1
        plugin.source.write_bytes(plugin.source.read_bytes().replace(b"Original", "Fresh 中文".encode()))
        image.write_bytes(png("blue"))
        result = await export_html(plugin, opened)
        effect, = result["hostEffects"]
        record = plugin.artifacts.get(bot_alias="main", artifact_id=effect["artifactId"])
        assert effect["type"] == "download_artifact" and record.filename == effect["filename"] == "sample.docx"
        assert record.content_type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        document = Document(record.path)
        assert "Fresh 中文" in [p.text for p in document.paragraphs]
        assert "Original" not in [p.text for p in document.paragraphs]
        assert len(document.inline_shapes) == 1
        assert any(part.blob == png("blue") for part in document.part.package.parts if part.content_type == "image/png")
        methods = [event["host_api_method"] for event in plugin.audit if event["event"] == "host_api"]
        assert methods.count("host.workspace.read_bytes") == 4
        assert methods.count("host.temp.write_artifact") == 1
        auth = AuthContext(user_id=1, token_used=True, capabilities={CAP_RUN_PLUGINS})
        manager = SimpleNamespace(main_profile=SimpleNamespace(alias="main"),
                                  managed_profiles={"other": SimpleNamespace(alias="other")},
                                  plugin_service=SimpleNamespace(get_artifact=plugin.artifacts.get))
        server = SimpleNamespace(manager=manager, _with_capability=AsyncMock(return_value=auth),
                                 _manager_alias=lambda request: request.match_info["alias"])
        request = make_mocked_request("GET", "/artifact", match_info={"alias": "main", "artifact_id": record.artifact_id})
        response = await WebApiServer.download_plugin_artifact(server, request)
        assert isinstance(response, web.FileResponse) and response._path == record.path
        assert response.headers["Content-Type"] == record.content_type
        assert response.headers["Content-Disposition"] == 'attachment; filename="sample.docx"'
        other = make_mocked_request("GET", "/artifact", match_info={"alias": "other", "artifact_id": record.artifact_id})
        with pytest.raises(WebApiError) as caught:
            await WebApiServer.download_plugin_artifact(server, other)
        assert caught.value.code == "plugin_artifact_not_found"


@pytest.mark.asyncio
async def test_parent_relative_image_stays_within_host_workspace(tmp_path, converter_available):
    async with copied_plugin(tmp_path) as plugin:
        (plugin.workspace / "reports").mkdir()
        (plugin.workspace / "assets").mkdir()
        (plugin.workspace / "assets/image.png").write_bytes(png("red"))
        (plugin.workspace / "reports/report.html").write_text('<p>正文</p><img src="../assets/image.png">', encoding="utf-8")
        opened = await open_html(plugin, "reports/report.html")
        assert not opened["summary"]["actions"][0]["disabled"]
        assert (await export_html(plugin, opened))["hostEffects"][0]["type"] == "download_artifact"


@pytest.mark.asyncio
@pytest.mark.parametrize("reference", ["missing.png", "../outside.png", "%2e%2e/outside.png",
                                     "https://example.invalid/x.png", "file:///C:/outside.png"])
async def test_missing_external_and_outside_images_block_changed_export(tmp_path, converter_available, reference):
    (tmp_path / "outside.png").write_bytes(png("red"))
    async with copied_plugin(tmp_path) as plugin:
        opened = await open_html(plugin)
        plugin.source.write_text(f'<p>Changed</p><img src="{reference}">', encoding="utf-8")
        result = await export_html(plugin, opened)
        assert result["message"] and not result.get("hostEffects")
        reopened = await open_html(plugin)
        assert reopened["summary"]["actions"][0]["disabled"]
        assert reopened["initialWindow"]["rows"][0]["cells"]["errors"]
        assert not plugin.artifacts._records


@pytest.mark.asyncio
@pytest.mark.parametrize("failure", ["not-installed", "import-failure"])
async def test_missing_dependency_keeps_view_available_with_host_python_command(tmp_path, failure):
    core = "raise ModuleNotFoundError(\"No module named 'docx'\", name='docx')" if failure == "import-failure" else None
    async with copied_plugin(tmp_path, missing=failure == "not-installed", core=core) as plugin:
        opened = await open_html(plugin)
        errors = opened["initialWindow"]["rows"][0]["cells"]["errors"]
        assert "缺少转换依赖" in errors
        assert sys.executable in errors and str(plugin.root / "requirements.txt") in errors
        assert "-m pip install -r" in errors and opened["summary"]["actions"][0]["disabled"]
        result = await export_html(plugin, opened)
        assert "缺少转换依赖" in result["message"] and not result.get("hostEffects")


@pytest.mark.asyncio
@pytest.mark.parametrize("limit", [{"maxSourceBytes": 10}, {"readBytes": 10}])
async def test_source_limits_are_checked_before_read(tmp_path, limit):
    async with copied_plugin(tmp_path, core=STUB_CORE,
                             config=limit if "maxSourceBytes" in limit else None,
                             limits=limit if "readBytes" in limit else None) as plugin:
        opened = await open_html(plugin)
        assert "限额" in opened["initialWindow"]["rows"][0]["cells"]["errors"]
        assert not any(event.get("host_api_method") == "host.workspace.read_bytes" for event in plugin.audit)


@pytest.mark.asyncio
async def test_resource_and_artifact_limits_do_not_publish(tmp_path):
    core = STUB_CORE.replace('"resources": []', '"resources": ["large.png"]')
    async with copied_plugin(tmp_path / "resources", core=core, config={"maxResourceBytes": 3}) as plugin:
        (plugin.workspace / "large.png").write_bytes(b"large-image")
        opened = await open_html(plugin)
        assert "限额" in opened["initialWindow"]["rows"][0]["cells"]["errors"]
        assert not plugin.artifacts._records
    async with copied_plugin(tmp_path / "artifact", core=STUB_CORE, limits={"artifactBytes": 8}) as plugin:
        result = await export_html(plugin, await open_html(plugin))
        assert "宿主限额" in result["message"] and not result.get("hostEffects")


@pytest.mark.asyncio
@pytest.mark.parametrize("stage", ["inspect", "export"])
async def test_worker_timeout_keeps_runtime_responsive(tmp_path, stage):
    core = STUB_CORE + f'\nimport time\ndef {"inspect_html" if stage == "inspect" else "export_docx"}(source, resources=None, config=None):\n    time.sleep(30)\n'
    async with copied_plugin(tmp_path, core=core, config={"conversionTimeoutSeconds": 1}) as plugin:
        opened = await open_html(plugin)
        result = opened["initialWindow"]["rows"][0]["cells"]["errors"] if stage == "inspect" else (await export_html(plugin, opened))["message"]
        assert "转换超时: 1s" in result
        assert plugin.runtime.active_process_count() == 1
        window = await plugin.runtime.get_view_window("main", plugin.manifest, opened["sessionId"], {})
        assert window["totalRows"] == 1 and not plugin.artifacts._records


@pytest.mark.asyncio
async def test_host_read_and_parse_share_one_deadline_and_late_reply_is_ignored(tmp_path):
    class SlowHost(PluginHostApi):
        async def dispatch(self, context, manifest, method, params):
            if method == "host.workspace.read_bytes":
                await asyncio.sleep(1.4)
            return await super().dispatch(context, manifest, method, params)

    async with copied_plugin(tmp_path, core=STUB_CORE, config={"conversionTimeoutSeconds": 1}, host_type=SlowHost) as plugin:
        opened = await open_html(plugin)
        assert "转换超时: 1s" in opened["initialWindow"]["rows"][0]["cells"]["errors"]
        await asyncio.sleep(0.5)
        window = await plugin.runtime.get_view_window("main", plugin.manifest, opened["sessionId"], {})
        assert window["totalRows"] == 1
        assert plugin.runtime.active_process_count() == 1


@pytest.mark.asyncio
async def test_conversion_stages_cannot_each_take_a_full_timeout(tmp_path):
    core = STUB_CORE.replace('    return {"title"', '    import time\n    time.sleep(0.6)\n    return {"title"')
    async with copied_plugin(tmp_path, core=core, config={"conversionTimeoutSeconds": 1}) as plugin:
        opened = await open_html(plugin)
        assert "转换超时: 1s" in opened["initialWindow"]["rows"][0]["cells"]["errors"]
        assert not plugin.artifacts._records


def test_worker_receives_remaining_budget_host_python_and_clamped_configuration(monkeypatch):
    spec = importlib.util.spec_from_file_location("html_docx_runner", PLUGIN / "backend/worker_runner.py")
    runner = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(runner)
    launch = Mock(side_effect=subprocess.TimeoutExpired("worker", 0.5))
    monkeypatch.setattr(runner.subprocess, "run", launch)
    config = runner.bounded_config({"conversionTimeoutSeconds": 999, "maxImagePixels": 99999999})
    assert config["conversionTimeoutSeconds"] == 55 and config["maxImagePixels"] == 16000000
    with pytest.raises(TimeoutError):
        runner.run_worker("inspect", HTML, config, deadline=time.monotonic() + 0.5)
    assert launch.call_args.args[0] == [sys.executable, str(PLUGIN / "backend/worker.py")]
    assert 0 < launch.call_args.kwargs["timeout"] <= 0.5
