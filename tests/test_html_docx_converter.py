from __future__ import annotations

import base64
import importlib.util
from io import BytesIO
from pathlib import Path

import pytest

for dependency in ("lxml", "tinycss2", "cssselect2", "docx", "PIL", "resvg_py"):
    pytest.importorskip(dependency)

from docx import Document
from docx.enum.text import WD_LINE_SPACING
from docx.oxml.ns import qn
from PIL import Image


BACKEND = Path(__file__).resolve().parents[1] / "examples/plugins/html-docx/backend"


@pytest.fixture(scope="module")
def converter():
    spec = importlib.util.spec_from_file_location("tested_html_docx_converter", BACKEND / "converter.py")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def html(body):
    return f"<!doctype html><html><head><meta charset='utf-8'><title>中文报告 Report</title></head><body>{body}</body></html>".encode()


def png(width=80, height=40):
    output = BytesIO()
    Image.new("RGB", (width, height), "red").save(output, format="PNG")
    return output.getvalue()


def exported(converter, source, resources=None, config=None):
    data, report = converter.export_docx(source, resources, config)
    assert report["errors"] == [], report
    assert data.startswith(b"PK")
    return Document(BytesIO(data)), report


def test_native_text_order_chinese_whitespace_links_and_a4(converter):
    document, report = exported(converter, html('''
        <h1 id="top">标题 Title</h1><div>前言 <strong>加粗</strong> 和 <em>斜体</em>。
        <p>正文 <a href="https://example.org/report">链接 中文</a> <u>下划线</u> x<sup>2</sup> H<sub>2</sub>O&nbsp;尾</p>
        <blockquote><p>引用</p></blockquote><pre>  代码\n\tprint("中文")  </pre>
        <p><a href="#top">回到标题</a><br>第二行</p>末尾</div>'''))
    assert [p.text for p in document.paragraphs] == ["标题 Title", "前言 加粗 和 斜体。", "正文 链接 中文 下划线 x2 H2O\xa0尾", "引用", '  代码\n\tprint("中文")  ', "回到标题\n第二行", "末尾"]
    assert report["paragraphCount"] == 7 and report["title"] == "中文报告 Report"
    assert document.paragraphs[0].style.name == "Heading 1"
    assert document.paragraphs[0].paragraph_format.keep_with_next
    assert document.paragraphs[3].paragraph_format.left_indent.pt == 24
    section = document.sections[0]
    assert section.page_width.mm == pytest.approx(210, abs=.05)
    assert section.page_height.mm == pytest.approx(297, abs=.05)
    assert section.left_margin.mm == pytest.approx(20, abs=.05)
    assert document.paragraphs[2]._p.xpath(".//w:hyperlink[@r:id]")
    assert document.paragraphs[5]._p.xpath(".//w:hyperlink[@w:anchor='html_1']")
    assert document.paragraphs[2]._p.xpath(".//w:vertAlign[@w:val='superscript']")
    assert document.paragraphs[2]._p.xpath(".//w:vertAlign[@w:val='subscript']")
    assert all(run._r.rPr.rFonts.get(qn("w:eastAsia")) for paragraph in document.paragraphs for run in paragraph.runs if run.text)
    document.paragraphs[0].runs[0].text = "可编辑"
    saved = BytesIO()
    document.save(saved)
    assert Document(BytesIO(saved.getvalue())).paragraphs[0].text == "可编辑"


def test_css_cascade_inheritance_important_and_units(converter):
    document, report = exported(converter, html('''<style>
      body {font-family: "Microsoft YaHei"; font-size: 12pt; color: #123456}
      p {font-size: 1.5em; line-height: 1.6; margin: 4pt 8pt 10pt 12pt}
      .accent {color: red !important} #specific {color: green}
      p > span {font-weight: 700; font-size: 50%}
    </style><p style="text-align:center">甲<span class="accent" id="specific" style="color:blue">乙</span><span class="accent" style="color: blue !important">丙</span></p>'''))
    paragraph, = document.paragraphs
    assert [(r.text, r.font.size.pt, str(r.font.color.rgb), r.font.bold) for r in paragraph.runs] == [
        ("甲", 18, "123456", False), ("乙", 9, "FF0000", True), ("丙", 9, "0000FF", True)]
    assert paragraph.paragraph_format.line_spacing.pt == pytest.approx(28.8)
    assert paragraph.paragraph_format.line_spacing_rule == WD_LINE_SPACING.AT_LEAST
    assert paragraph.paragraph_format.space_after.pt == 10
    assert paragraph.paragraph_format.left_indent.pt == 12
    assert paragraph.paragraph_format.right_indent.pt == 8
    assert report["warnings"] == []


def test_native_lists_nested_start_and_restart_preserve_content(converter):
    document, report = exported(converter, html('''<ol start="3" type="A"><li><p>第一项</p><p>继续段落</p>
      <ul><li>嵌套 <b>强调</b></li><li>另一个</li></ul></li><li value="8">第八项</li><li>第九项</li></ol>'''))
    assert [p.text for p in document.paragraphs] == ["第一项", "继续段落", "嵌套 强调", "另一个", "第八项", "第九项"]
    numbered = [(p.text, p._p.pPr.numPr.ilvl.val, p._p.pPr.numPr.numId.val) for p in document.paragraphs if p._p.pPr.numPr is not None]
    assert [item[:2] for item in numbered] == [("第一项", 0), ("嵌套 强调", 1), ("另一个", 1), ("第八项", 0), ("第九项", 0)]
    assert numbered[-1][2] == numbered[-2][2] != numbered[0][2]
    xml = document.part.numbering_part.element
    assert xml.xpath(".//w:numFmt[@w:val='upperLetter']")
    assert xml.xpath(".//w:startOverride[@w:val='3']") and xml.xpath(".//w:startOverride[@w:val='8']")


def test_native_merged_table_headers_width_borders_shading_alignment(converter):
    document, report = exported(converter, html('''<style>td,th {border: 1pt solid #345678;padding:3pt} td {vertical-align: middle}</style>
      <table style="width:100%"><colgroup><col style="width:30%"><col style="width:70%"></colgroup>
      <thead><tr><th colspan="2" style="background:#eee">表头</th></tr></thead><tbody>
      <tr><td rowspan="2">合并</td><td style="text-align:right">数值 1</td></tr><tr><td>数值 2</td></tr></tbody></table>'''))
    table, = document.tables
    assert report["tableCount"] == 1
    assert table.cell(0, 0)._tc is table.cell(0, 1)._tc
    assert table.cell(1, 0)._tc is table.cell(2, 0)._tc
    assert table.cell(0, 0).text == "表头" and table.cell(1, 0).text == "合并"
    assert table.rows[0]._tr.xpath("./w:trPr/w:tblHeader")
    assert table.cell(0, 0)._tc.xpath("./w:tcPr/w:shd[@w:fill='EEEEEE']")
    assert not table.cell(0, 0)._tc.xpath("./w:p/w:pPr/w:shd|./w:p/w:r/w:rPr/w:shd")
    assert table.cell(1, 0)._tc.xpath("./w:tcPr/w:tcBorders/w:top[@w:color='345678']")
    assert not table.cell(1, 0)._tc.xpath("./w:p/w:pPr/w:pBdr")
    assert table.cell(1, 0)._tc.xpath("./w:tcPr/w:vAlign[@w:val='center']")
    assert table.cell(1, 1).paragraphs[0]._p.xpath("./w:pPr/w:jc[@w:val='right']")
    assert sum(c.width.pt for c in table.columns) == pytest.approx(170 * 72 / 25.4, abs=.1)
    assert table.columns[0].width.pt / table.columns[1].width.pt == pytest.approx(3 / 7, abs=.001)


def test_images_discovery_validation_raw_paths_and_fit(converter):
    source = html('<p>前<img src="images/中文 image.png" width="2000" alt="说明">后</p>')
    discovery = converter.inspect_html(source)
    assert discovery["resources"] == ["images/中文 image.png"] and discovery["errors"] == []
    assert converter.inspect_html(source, {})["errors"]
    data, missing = converter.export_docx(source)
    assert data == b"" and any("Missing image" in error for error in missing["errors"])
    assert converter.inspect_html(source, {"images/中文 image.png": b"not an image"})["errors"]
    document, report = exported(converter, source, {"images/中文 image.png": png()})
    shape, = document.inline_shapes
    assert shape.width.pt == pytest.approx(170 * 72 / 25.4, abs=.1)
    assert shape.width / shape.height == pytest.approx(2)
    assert report["imageCount"] == 1
    assert shape._inline.docPr.get("descr") == "说明"


def test_base64_png_jpeg_and_inline_svg_static_chart(converter):
    jpeg = BytesIO()
    Image.new("RGB", (10, 10), "blue").save(jpeg, "JPEG")
    source = html(f'''<p><img src="data:image/png;base64,{base64.b64encode(png()).decode()}">
      <img src="data:image/jpeg;base64,{base64.b64encode(jpeg.getvalue()).decode()}"></p>
      <svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 160 80" width="160" height="80">
        <defs><linearGradient id="g"><stop offset="0" stop-color="red"/><stop offset="1" stop-color="blue"/></linearGradient></defs>
        <rect width="160" height="80" fill="url(#g)"/><text x="5" y="35" fill="white">中文 Chart</text></svg>''')
    document, report = exported(converter, source)
    assert len(document.inline_shapes) == report["imageCount"] == 3
    assert report["resources"] == []
    pictures = [rel.target_part.blob for rel in document.part.rels.values() if rel.reltype.endswith("/image")]
    chart = Image.open(BytesIO(pictures[-1])).convert("RGB")
    assert chart.size == (160, 80)
    assert chart.getpixel((1, 70))[0] > 240 and chart.getpixel((158, 70))[2] > 240


def test_page_breaks_hidden_content_and_located_simplification_warnings(converter):
    document, report = exported(converter, html('''<div id="layout" style="display:grid;box-shadow:0 0 4px #000">
      <h2 style="break-before:page">新页</h2><p style="page-break-after:always;page-break-inside:avoid">正文</p>
      <p hidden>隐藏</p><p style="display:none"><img src="missing.png"></p></div>'''))
    assert [p.text for p in document.paragraphs] == ["新页", "正文"]
    assert document.paragraphs[0].paragraph_format.page_break_before
    assert document.paragraphs[1].paragraph_format.keep_together
    assert document.paragraphs[1]._p.xpath(".//w:br[@w:type='page']")
    assert all("#layout" in warning and "line" in warning for warning in report["warnings"])
    assert len(report["warnings"]) == 2 and report["imageCount"] == 0


@pytest.mark.parametrize("content", [
    '<script>document.write("bad")</script>', '<canvas id="chart"></canvas>', '<p onclick="evil()">x</p>',
    '<iframe src="https://example.org"></iframe>', '<link rel="stylesheet" href="https://example.org/style.css">',
    '<a href="javascript:evil()">bad</a>', '<table><tr><td>one</td><td>two</td></tr><tr><td>missing</td></tr></table>',
    '<table><tr><td colspan="999999999">bad</td></tr></table>', '<table><tr><td rowspan="2">bad</td></tr></table>',
    '<ul>lost text<li>item</li></ul>', '<ol reversed><li>item</li></ol>', '<td>orphan</td>',
])
def test_unreliable_or_dynamic_content_blocks_export_with_location(converter, content):
    data, report = converter.export_docx(html(content))
    assert data == b"" and report["errors"]
    assert any("line" in message for message in report["errors"])


@pytest.mark.parametrize("src", ["https://example.org/image.png", "//server/image.png", "file:///secret.png", "C:\\secret.png", "/secret.png", "image.png?size=10", "image.png#fragment", "data:image/svg+xml;base64,AAAA"])
def test_image_references_cannot_escape_host_resource_contract(converter, src):
    data, report = converter.export_docx(html(f'<img src="{src}">'), {src: png()})
    assert data == b"" and report["errors"] and report["resources"] == []


@pytest.mark.parametrize("body", [
    '<image href="file:///secret.png"/>', '<use xlink:href="https://example.org/x.svg#id"/>',
    '<rect style="fill:url(https://example.org/x.svg)"/>', '<rect fill="url(&quot;file:///secret&quot;)"/>',
    '<style>@import "file:///secret";</style>', '<style>rect{fill:u\\72l(\\66ile:///secret)}</style>',
    '<foreignObject><p>unreliable</p></foreignObject>', '<script>alert(1)</script>',
    '<animate attributeName="x"/>', '<image href="data:image/svg+xml;base64,AAAA"/>',
    '<use href="#missing"/>', '<g id="self"><use href="#self"/></g>',
    '<defs><g id="a"><use href="#b"/></g><g id="b"><use href="#a"/></g></defs><use href="#a"/>',
])
def test_svg_never_passes_unsafe_references_to_renderer(converter, monkeypatch, body):
    import resvg_py
    calls = []
    monkeypatch.setattr(resvg_py, "svg_to_bytes", lambda **kwargs: calls.append(kwargs))
    data, report = converter.export_docx(html(f'<svg width="100" height="80">{body}</svg>'))
    assert data == b"" and report["errors"] and calls == []


def test_security_limits_entities_depth_and_configuration_clamping(converter):
    cases = [
        (b'<!DOCTYPE html [<!ENTITY x SYSTEM "file:///secret">]><p>&x;</p>', {}),
        (b"\xff", {}), (b"<p>\x00</p>", {}),
        (html("<div>" * 66 + "x" + "</div>" * 66), {}),
        (html("<p>one</p><p>two</p>"), {"maxElements": 3}),
        (html("<p>hello</p>"), {"maxSourceBytes": 10}),
        (html("<table><tr><td>1</td><td>2</td></tr></table>"), {"maxTableCells": 1}),
    ]
    for source, config in cases:
        data, report = converter.export_docx(source, config=config)
        assert data == b"" and report["errors"], (source, report)
    source = html('<img src="x.png">')
    assert converter.inspect_html(source, {"x.png": png()}, {"maxImagePixels": 100})["errors"]
    assert converter.inspect_html(source, {"x.png": png()}, {"maxResourceBytes": 10})["errors"]
    assert converter._limit({"maxElements": 10**20}, "maxElements") == 10000
    document, _ = exported(converter, html("<p>x</p>"), config={"defaultFontSizePt": float("nan"), "maxElements": "invalid"})
    assert document.paragraphs[0].runs[0].font.size.pt == 11


def test_export_uses_no_source_resource_filesystem_or_network_io(converter, monkeypatch):
    import builtins
    import socket
    original = builtins.open
    opened = []
    def guarded(file, *args, **kwargs):
        opened.append(str(file))
        # python-docx opens its bundled default template and Pillow loads plugins.
        assert "private.png" not in str(file)
        return original(file, *args, **kwargs)
    monkeypatch.setattr(builtins, "open", guarded)
    monkeypatch.setattr(socket, "socket", lambda *args, **kwargs: pytest.fail("network access"))
    exported(converter, html('<img src="private.png">'), {"private.png": png()})
    assert "private.png" not in opened


def test_empty_page_break_container_margins_and_list_block_children(converter):
    document, report = exported(converter, html('''<p>Before</p><div style="page-break-before:always"></div>
      <div style="margin-top:20pt;margin-bottom:15pt"><p>A</p><p>B</p></div>
      <ol><li><div>One</div><div>Two</div></li><li><p>Three</p></li></ol>'''))
    assert [p.text for p in document.paragraphs] == ["Before", "", "A", "B", "One", "Two", "Three"]
    assert document.paragraphs[1].paragraph_format.page_break_before
    assert document.paragraphs[2].paragraph_format.space_before.pt == 20
    assert document.paragraphs[3].paragraph_format.space_after.pt == 15
    assert document.paragraphs[4]._p.pPr.numPr is not None
    assert document.paragraphs[5]._p.pPr.numPr is None
    assert document.paragraphs[6]._p.pPr.numPr is not None


def test_svg_fragment_after_styles_and_outer_css_are_preserved(converter):
    document, report = exported(converter, b'<style>.bar{fill:red}</style><svg width="100" height="50"><rect class="bar" width="100" height="50"/></svg>')
    assert len(document.inline_shapes) == 1
    blob = next(rel.target_part.blob for rel in document.part.rels.values() if rel.reltype.endswith("/image"))
    assert Image.open(BytesIO(blob)).convert("RGB").getpixel((50, 25)) == (255, 0, 0)
    assert report["paragraphCount"] == 1


def test_svg_external_url_from_html_css_is_rejected_before_render(converter, monkeypatch):
    import resvg_py
    monkeypatch.setattr(resvg_py, "svg_to_bytes", lambda **kwargs: pytest.fail("unsafe SVG renderer call"))
    data, report = converter.export_docx(html('<style>.bar{fill:url(file:///private)}</style><svg width="100" height="50"><rect class="bar" width="100" height="50"/></svg>'))
    assert data == b"" and any("external" in error for error in report["errors"])


def test_nested_tables_and_rowspan_zero_retain_all_cell_content(converter):
    document, report = exported(converter, html('''<table><tbody><tr><td rowspan="0">合并</td><td>before<table><tr><td>inner</td></tr></table>after</td></tr><tr><td>last</td></tr></tbody></table>'''))
    table, = document.tables
    assert table.cell(0, 0)._tc is table.cell(1, 0)._tc
    cell = table.cell(0, 1)
    assert [p.text for p in cell.paragraphs if p.text] == ["before", "after"]
    assert cell.tables[0].cell(0, 0).text == "inner"
    assert report["tableCount"] == 2 and report["warnings"] == []


def test_inherited_indent_and_nested_superscript_are_native(converter):
    document, report = exported(converter, html('<div style="text-indent:2em"><p>公式 x<sup><b>2</b></sup></p></div>'))
    paragraph, = document.paragraphs
    assert paragraph.paragraph_format.first_line_indent.pt == 22
    assert paragraph.runs[-1].font.superscript and paragraph.runs[-1].font.bold


def test_svg_pixel_limit_and_combined_inline_resource_limit(converter):
    data, report = converter.export_docx(html('<svg width="100000" height="100000"><rect width="5" height="5"/></svg>'))
    assert data == b"" and any("maxImagePixels" in error for error in report["errors"])
    picture = png()
    encoded = base64.b64encode(picture).decode()
    data, report = converter.export_docx(html(f'<img src="x.png"><img src="data:image/png;base64,{encoded}">'), {"x.png": picture}, {"maxResourceBytes": len(picture) + 1})
    assert data == b"" and any("maxResourceBytes" in error for error in report["errors"])


def test_merged_header_width_does_not_override_individual_columns(converter):
    document, report = exported(converter, html('''<p>Before</p><table style="width:400pt;margin:10pt 0 12pt">
      <thead><tr><th colspan="2" style="width:400pt">Group</th></tr><tr><th style="width:100pt">A</th><th style="width:300pt">B</th></tr></thead>
      <tbody><tr><td>One</td><td>Two</td></tr></tbody></table><p>After</p>'''))
    assert [column.width.pt for column in document.tables[0].columns] == [100, 300]
    assert document.paragraphs[0].paragraph_format.space_after.pt == 10
    assert document.paragraphs[1].paragraph_format.space_before.pt == 12


def test_css_line_height_uses_points_without_grid_and_expands_for_larger_runs_and_images(converter):
    document, report = exported(converter, html('''<style>body{font-size:11pt;line-height:1.35}</style>
      <p>中文正文<br>第二行</p><p>Text <span style="font-size:22pt">大字</span></p>
      <p>Image <img src="x.png"></p><div id="fixed"><p style="line-height:16pt">Explicit height</p></div>
      <p style="line-height:16pt"><img src="x.png"></p>
      <p style="line-height:16pt"><span style="font-size:30pt">Large text</span></p>'''), {"x.png": png(80, 100)})
    assert not document.sections[0]._sectPr.xpath("./w:docGrid")
    assert document.styles["Normal"].element.xpath("./w:pPr/w:snapToGrid[@w:val='0']")
    plain, mixed, image, fixed, fixed_image, fixed_large = document.paragraphs
    assert plain.paragraph_format.line_spacing.pt == pytest.approx(14.85)
    assert plain._p.xpath("./w:pPr/w:rPr/w:sz[@w:val='22']")
    assert mixed.paragraph_format.line_spacing.pt == pytest.approx(29.7)
    for paragraph in (plain, mixed, image, fixed_image, fixed_large):
        assert paragraph.paragraph_format.line_spacing_rule == WD_LINE_SPACING.AT_LEAST
    assert fixed.paragraph_format.line_spacing_rule == WD_LINE_SPACING.EXACTLY
    assert fixed.paragraph_format.line_spacing.pt == 16
    assert all(shape.height.pt == 75 for shape in document.inline_shapes)


def test_cell_box_formatting_does_not_duplicate_or_remove_explicit_child_formatting(converter):
    document, report = exported(converter, html('''<table><tr><td style="border:1pt solid red;background:#eee;padding:4pt">
      direct <span style="background:yellow">highlight</span> tail
      <p style="border:1pt solid blue;background:#abc">boxed paragraph</p>
    </td></tr></table>'''))
    cell = document.tables[0].cell(0, 0)
    assert cell._tc.xpath("./w:tcPr/w:tcBorders/w:top[@w:color='FF0000']")
    assert cell._tc.xpath("./w:tcPr/w:shd[@w:fill='EEEEEE']")
    assert cell._tc.xpath("./w:tcPr/w:tcMar/w:left[@w:w='80']")
    direct, boxed = cell.paragraphs
    assert not direct._p.xpath("./w:pPr/w:pBdr|./w:pPr/w:shd")
    fills = direct._p.xpath(".//w:rPr/w:shd/@w:fill")
    assert fills == ["FFFF00"]
    assert boxed._p.xpath("./w:pPr/w:pBdr/w:top[@w:color='0000FF']")
    assert boxed._p.xpath("./w:pPr/w:shd[@w:fill='AABBCC']")
