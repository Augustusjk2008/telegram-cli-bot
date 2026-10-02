from __future__ import annotations

import importlib.util
from io import BytesIO
from pathlib import Path

import pytest

for dependency in ("pptx", "resvg_py", "PIL", "tinycss2", "cssselect2", "defusedxml"):
    pytest.importorskip(dependency)

from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.enum.text import MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn


BACKEND = Path(__file__).resolve().parents[1] / "examples/plugins/svg-pptx/backend"


@pytest.fixture
def converter(monkeypatch):
    monkeypatch.syspath_prepend(str(BACKEND))
    spec = importlib.util.spec_from_file_location("tested_svg_converter", BACKEND / "converter.py")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def svg(body: str, attributes: str = 'width="800" height="450"') -> bytes:
    return f'<svg xmlns="http://www.w3.org/2000/svg" {attributes}>{body}</svg>'.encode("utf-8")


def test_bilingual_rich_text_inherits_css_units_and_viewbox_and_remains_editable(converter):
    source = svg('''
      <style>g {font-family: "Microsoft YaHei", sans-serif; font-size: 24pt; fill: #123456}
        .accent {fill: red !important; font-size: 1.25em}
      </style>
      <g font-weight="bold"><text x="140" y="300">中文 <tspan class="accent" style="fill: blue">Hello</tspan>!</text></g>
    ''', 'width="10in" height="5.625in" viewBox="100 200 480 270"')
    content, report = converter.export_pptx(source)
    assert report["errors"] == []
    presentation = Presentation(BytesIO(content))
    assert len(presentation.slides) == 1
    assert presentation.slide_width / presentation.slide_height == pytest.approx(16 / 9)
    shape, = presentation.slides[0].shapes
    assert shape.text == "中文 Hello!"
    assert shape.left == 80 * 9525
    runs = shape.text_frame.paragraphs[0].runs
    assert [(run.text, run.font.size.pt, str(run.font.color.rgb), run.font.bold) for run in runs] == [
        ("中文 ", 48, "123456", True), ("Hello", 60, "FF0000", True), ("!", 48, "123456", True),
    ]
    assert all(run._r.rPr.find(qn("a:ea")).get("typeface") == run.font.name for run in runs)
    assert shape.text_frame.word_wrap is False and shape.text_frame.auto_size == MSO_AUTO_SIZE.NONE
    assert shape.text_frame.margin_top == shape.text_frame.margin_left == shape.text_frame.margin_right == shape.text_frame.margin_bottom == 0
    assert shape.text_frame.paragraphs[0].line_spacing == shape.height
    runs[0].text = "修改后的中文 "
    runs[1].text = "Edited English"
    saved = BytesIO()
    presentation.save(saved)
    assert Presentation(BytesIO(saved.getvalue())).slides[0].shapes[0].text == "修改后的中文 Edited English!"


def test_positioned_tspans_split_lines_and_continuous_spans_share_box(converter):
    source = svg('''<text x="200" y="100" text-anchor="middle" font-family="Microsoft YaHei" font-size="28">
      <tspan>第一行 <tspan fill="red">中文</tspan></tspan>
      <tspan x="200" dy="40">Second line</tspan>
    </text>''')
    content, report = converter.export_pptx(source)
    assert report["textCount"] == 2
    first, second = Presentation(BytesIO(content)).slides[0].shapes
    assert (first.text, second.text) == ("第一行 中文", "Second line")
    assert second.top - first.top == 40 * 9525
    for shape in (first, second):
        assert shape.left + shape.width / 2 == pytest.approx(200 * 9525, abs=1)
        assert shape.text_frame.paragraphs[0].alignment == PP_ALIGN.CENTER


def test_nested_translation_rotation_and_uniform_scale_preserve_native_text(converter):
    source = svg('<g transform="translate(50 20)"><g transform="rotate(30 100 100) scale(2)">'
                 '<text x="50" y="50" font-family="Arial" font-size="20">Rotated</text></g></g>')
    content, report = converter.export_pptx(source)
    assert report["errors"] == []
    shape, = Presentation(BytesIO(content)).slides[0].shapes
    assert shape.rotation == pytest.approx(30)
    assert shape.text_frame.paragraphs[0].runs[0].font.size.pt == 30
    # The transformed baseline is (150, 120). Recover it from the rotated box.
    import math
    half_height = (shape.height / 9525) / 2
    import fonts
    face = fonts.resolve("Arial", False, False, "Rotated", [])[0][1]
    _, ascent, _, _ = fonts.measure(face, 20, "Rotated")
    local_y = ascent * 2 - half_height
    local_x = -shape.width / 9525 / 2
    baseline_x = (shape.left + shape.width / 2) / 9525 + local_x * math.cos(math.pi / 6) - local_y * math.sin(math.pi / 6)
    baseline_y = (shape.top + shape.height / 2) / 9525 + local_x * math.sin(math.pi / 6) + local_y * math.cos(math.pi / 6)
    assert (baseline_x, baseline_y) == pytest.approx((150, 120), abs=0.001)


def test_graphics_layers_have_no_text_and_keep_occlusion_and_css_selector_colors(converter):
    source = svg('''<style>g > rect:nth-child(3) { fill: red }</style><g>
      <rect width="800" height="450" fill="white"/>
      <text x="30" y="80" font-size="32">Covered text</text>
      <rect x="20" y="50" width="200" height="45"/>
      <text x="300" y="80">Foreground</text></g>''')
    content, report = converter.export_pptx(source, {"renderMaxSide": 800})
    assert (report["textCount"], report["layerCount"]) == (2, 2)
    shapes = Presentation(BytesIO(content)).slides[0].shapes
    assert [shape.shape_type for shape in shapes] == [
        MSO_SHAPE_TYPE.PICTURE, MSO_SHAPE_TYPE.TEXT_BOX, MSO_SHAPE_TYPE.PICTURE, MSO_SHAPE_TYPE.TEXT_BOX,
    ]
    background = Image.open(BytesIO(shapes[0].image.blob)).convert("RGBA")
    assert background.getextrema() == ((255, 255), (255, 255), (255, 255), (255, 255))
    foreground = Image.open(BytesIO(shapes[2].image.blob)).convert("RGBA")
    assert foreground.getpixel((40, 70)) == (255, 0, 0, 255)
    assert foreground.getpixel((320, 70))[3] == 0
    import svg_parser
    scene = svg_parser.parse(source)
    for layer in (item for item in scene.items if isinstance(item, svg_parser.Layer)):
        assert "<text" not in converter.graphics_svg(scene, layer)


@pytest.mark.parametrize(("body", "reason"), [
    ('<defs><path id="p" d="M0 0L100 0"/></defs><text id="bad"><textPath href="#p">Text</textPath></text>', "textPath"),
    ('<defs><mask id="m"><rect width="800" height="450" fill="white"/></mask></defs><text id="bad" mask="url(#m)">Text</text>', "mask"),
    ('<g transform="skewX(12)"><text id="bad">Text</text></g>', "倾斜"),
    ('<text id="bad" transform="scale(2 1)">Text</text>', "非等比"),
    ('<text id="bad" x="10 20">Text</text>', "单个位置"),
    ('<text id="bad" letter-spacing="2">Text</text>', "letter-spacing"),
    ('<defs><g id="ref"><text id="bad">Text</text></g></defs><use href="#ref"/>', "定义"),
    ('<g opacity="0.5"><text id="bad">Text</text><rect width="100" height="100"/></g>', "组合透明度"),
    ('<switch><text id="bad">Text</text><text>Other</text></switch>', "switch"),
    ('<flowRoot id="bad"><flowRegion><rect width="200" height="80"/></flowRegion><flowPara>Editable words</flowPara></flowRoot>', "流式文字"),
])
def test_unsupported_editable_text_has_element_reason_and_no_export(converter, body, reason):
    source = svg(body)
    report = converter.inspect_svg(source)
    assert any("#bad" in error and reason in error for error in report["errors"])
    with pytest.raises(ValueError, match=reason):
        converter.export_pptx(source)


@pytest.mark.parametrize("body", [
    '<image href="https://example.invalid/a.png"/>',
    '<image href="file:///C:/outside.png"/>',
    '<image href="data:image/svg+xml;base64,PHN2Zy8+"/>',
    '<rect style="fill: url(\'https://example.invalid/paint.svg\')"/>',
    '<style>rect{fill:u\\72l("https://example.invalid/paint.svg")}</style><rect/>',
    '<style>@import "https://example.invalid/style.css";</style>',
    '<script>alert(1)</script>',
    '<g onclick="alert(1)"><text>Text</text></g>',
    '<g xml:base="file:///C:/"><rect/></g>',
    '<defs><g id="cycle"><use href="#cycle"/></g></defs><use href="#cycle"/>',
])
def test_untrusted_resources_are_rejected_before_rasterization(converter, monkeypatch, body):
    import resvg_py
    render = pytest.fail
    monkeypatch.setattr(resvg_py, "svg_to_bytes", lambda **_: render("unsafe input reached renderer"))
    source = svg(body)
    assert converter.inspect_svg(source)["errors"]
    with pytest.raises(ValueError):
        converter.export_pptx(source)


@pytest.mark.parametrize("source", [
    b'<!DOCTYPE svg [<!ENTITY external SYSTEM "file:///C:/outside">]><svg xmlns="http://www.w3.org/2000/svg" width="800" height="450"><text>&external;</text></svg>',
    b'<?xml-stylesheet href="https://example.invalid/style.css"?><svg xmlns="http://www.w3.org/2000/svg" width="800" height="450"/>',
])
def test_xml_entities_and_processing_instructions_are_rejected(converter, source):
    assert converter.inspect_svg(source)["errors"]
    with pytest.raises(Exception):
        converter.export_pptx(source)


def test_missing_font_warns_and_preserves_text_in_local_fallback(converter):
    source = svg('<text x="20" y="80" font-family="Definitely Not Installed 123">中文 and English</text>')
    content, report = converter.export_pptx(source)
    assert any("Definitely Not Installed 123" in warning for warning in report["warnings"])
    shape, = Presentation(BytesIO(content)).slides[0].shapes
    assert shape.text == "中文 and English"
    assert all(run.font.name != "Definitely Not Installed 123" for run in shape.text_frame.paragraphs[0].runs)


def test_internal_graphical_references_and_embedded_raster_are_allowed(converter):
    import base64
    pixel = BytesIO()
    Image.new("RGB", (2, 2), "blue").save(pixel, "PNG")
    source = svg('<defs><rect id="r" width="40" height="40" fill="red"/></defs><use href="#r"/>'
                 f'<image href="data:image/png;base64,{base64.b64encode(pixel.getvalue()).decode()}" x="50" width="40" height="40"/>')
    content, report = converter.export_pptx(source, {"renderMaxSide": 800})
    assert report["errors"] == [] and report["textCount"] == 0
    picture, = Presentation(BytesIO(content)).slides[0].shapes
    image = Image.open(BytesIO(picture.image.blob)).convert("RGBA")
    assert image.getpixel((10, 10)) == (255, 0, 0, 255)
    assert image.getpixel((60, 10)) == (0, 0, 255, 255)


def test_use_keeps_referenced_graphics_from_other_layers_and_inherits_instance_fill(converter):
    source = svg('<rect id="painted" width="40" height="40" fill="red"/>'
                 '<text x="100" y="100">Editable</text>'
                 '<use href="#painted" x="100"/>'
                 '<defs><g id="group"><rect width="40" height="40"/></g></defs>'
                 '<use href="#group" x="200" fill="blue"/>')
    content, report = converter.export_pptx(source, {"renderMaxSide": 800})
    assert report["errors"] == []
    first, text, second = Presentation(BytesIO(content)).slides[0].shapes
    assert text.text == "Editable"
    before = Image.open(BytesIO(first.image.blob)).convert("RGBA")
    after = Image.open(BytesIO(second.image.blob)).convert("RGBA")
    assert before.getpixel((10, 10)) == (255, 0, 0, 255)
    assert after.getpixel((10, 10))[3] == 0
    assert after.getpixel((110, 10)) == (255, 0, 0, 255)
    assert after.getpixel((210, 10)) == (0, 0, 255, 255)


def test_parent_text_opacity_applies_to_continuous_tspans(converter):
    source = svg('<text x="20" y="80" opacity="0.25">A<tspan fill-opacity="0.5">B</tspan></text>')
    content, report = converter.export_pptx(source)
    assert report["errors"] == []
    shape, = Presentation(BytesIO(content)).slides[0].shapes
    alphas = [run._r.rPr.find(qn("a:solidFill")).find(qn("a:srgbClr")).find(qn("a:alpha")).get("val")
              for run in shape.text_frame.paragraphs[0].runs]
    assert alphas == ["25000", "12500"]


def test_center_anchor_includes_relative_dx_gap_and_inherited_preserved_spaces(converter):
    source = svg('<g xml:space="preserve"><text x="150" y="80" text-anchor="middle">A  B<tspan dx="20">C</tspan></text></g>')
    content, report = converter.export_pptx(source)
    assert report["errors"] == []
    first, second = Presentation(BytesIO(content)).slides[0].shapes
    assert (first.text, second.text) == ("A  B", "C")
    assert second.left - (first.left + first.width) == pytest.approx(20 * 9525, abs=1)
    assert (first.left + second.left + second.width) / 2 == pytest.approx(150 * 9525, abs=1)
