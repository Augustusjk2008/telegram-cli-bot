"""Static HTML to editable Word. Input resources are supplied exclusively by the host.

No document paths, resource resolvers, browser, or network client are used here.
resvg may use installed fonts for chart labels; all SVG resource references are
checked before it receives markup. The host runs this module in a timed worker.
"""
from __future__ import annotations

import base64
import binascii
from io import BytesIO
import math
import re
from urllib.parse import urlsplit, unquote

import cssselect2
from lxml import etree, html
from PIL import Image
import tinycss2
from tinycss2.color3 import parse_color
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.table import _Cell


LIMITS = {"maxSourceBytes": 2 * 1024**2, "maxResourceBytes": 16 * 1024**2,
          "maxImagePixels": 16_000_000, "maxElements": 10_000, "maxTableCells": 10_000}
MAX_DEPTH = 64
PAGE_WIDTH = 170 * 72 / 25.4
PAGE_HEIGHT = 257 * 72 / 25.4
SPACE = re.compile(r"[ \t\r\n\f]+")
NUMBER = re.compile(r"^([+-]?(?:\d*\.)?\d+(?:e[+-]?\d+)?)(pt|px|pc|in|cm|mm|em|rem|%)?$", re.I)
BLOCKS = set("html body div section article main header footer aside nav figure details summary dl blockquote".split())
PARAGRAPHS = set("p h1 h2 h3 h4 h5 h6 pre address dt dd figcaption".split())
INLINE = set("span b strong i em u s strike del ins sub sup code kbd samp var small big mark abbr cite q a br wbr time label".split())
IGNORED = {"head", "title", "meta", "style", "col", "colgroup"}
DYNAMIC = set("script canvas iframe frame frameset object embed video audio source track input select textarea button form math".split())
INHERITED = set("font-family font-size font-weight font-style color line-height text-align text-indent white-space text-decoration list-style-type visibility".split())
SUPPORTED = INHERITED | set("background-color vertical-align text-indent width height max-width max-height display float position break-before break-after break-inside page-break-before page-break-after page-break-inside border-collapse border-spacing table-layout".split())
for _side in ("top", "right", "bottom", "left"):
    SUPPORTED.update({f"margin-{_side}", f"padding-{_side}"})
    SUPPORTED.update(f"border-{_side}-{part}" for part in ("width", "style", "color"))
LIST_FORMATS = {"decimal": "decimal", "lower-alpha": "lowerLetter", "lower-latin": "lowerLetter",
                "upper-alpha": "upperLetter", "upper-latin": "upperLetter", "lower-roman": "lowerRoman",
                "upper-roman": "upperRoman", "disc": "bullet", "circle": "bullet", "square": "bullet", "none": "none"}


def _limit(config, name):
    try:
        value = int(config.get(name, LIMITS[name]))
        return min(LIMITS[name], max(1, value))
    except (TypeError, ValueError, OverflowError):
        return LIMITS[name]


def _length(value, font=11, reference=PAGE_WIDTH):
    match = NUMBER.fullmatch(str(value).strip())
    if not match:
        return None
    number = float(match[1])
    unit = (match[2] or "px").lower()
    result = number * {"pt": 1, "px": .75, "pc": 12, "in": 72, "cm": 72 / 2.54,
                       "mm": 72 / 25.4, "em": font, "rem": 11, "%": reference / 100}[unit]
    return result if math.isfinite(result) else None


def _color(value):
    try:
        color = parse_color(str(value).strip())
        if color is None or isinstance(color, str) or color.alpha == 0:
            return None
        return "".join(f"{round(255 * (component * color.alpha + 1 - color.alpha)):02X}"
                       for component in (color.red, color.green, color.blue))
    except (ValueError, TypeError):
        return None


def _element(parent, tag, **attributes):
    node = OxmlElement(tag)
    for key, value in attributes.items():
        node.set(qn("w:" + key), str(value))
    parent.append(node)
    return node


def _location(node):
    ident = f"#{node.get('id')}" if node.get("id") else ""
    return f"<{node.tag}{ident}> (line {node.sourceline or 1})"


def _declarations(content):
    """Expand the small set of Word-compatible shorthands before cascading."""
    result = []
    for decl in tinycss2.parse_declaration_list(content, skip_comments=True, skip_whitespace=True):
        if decl.type != "declaration":
            continue
        name, value = decl.lower_name, tinycss2.serialize(decl.value).strip()
        values = [tinycss2.serialize([token]) for token in decl.value if token.type != "whitespace"]
        expanded = {name: value}
        if name in {"margin", "padding", "border-width", "border-style", "border-color"} and 1 <= len(values) <= 4:
            parts = (values * 4)[:4] if len(values) == 1 else [values[0], values[1], values[0], values[1]]
            if len(values) == 3:
                parts = [values[0], values[1], values[2], values[1]]
            elif len(values) == 4:
                parts = values
            expanded = {}
            for side, part in zip(("top", "right", "bottom", "left"), parts):
                key = f"{name}-{side}" if name in {"margin", "padding"} else f"border-{side}-{name[7:]}"
                expanded[key] = part
        elif name == "border" or name in {"border-top", "border-right", "border-bottom", "border-left"}:
            border = {"width": "1px", "style": "none", "color": "black"}
            for part in values:
                if part in {"none", "hidden", "solid", "dotted", "dashed", "double", "groove", "ridge", "inset", "outset"}:
                    border["style"] = part
                elif _length(part) is not None or part in {"thin", "medium", "thick"}:
                    border["width"] = part
                elif _color(part):
                    border["color"] = part
            sides = ("top", "right", "bottom", "left") if name == "border" else (name[7:],)
            expanded = {f"border-{side}-{key}": val for side in sides for key, val in border.items()}
        elif name == "background" and (_color(value) or value == "transparent"):
            expanded = {"background-color": value}
        elif name == "text-decoration-line":
            expanded = {"text-decoration": value}
        result.extend((key, val, decl.important) for key, val in expanded.items())
    return result


class _Source:
    def __init__(self, source, resources, config):
        self.config = config if isinstance(config, dict) else {}
        self.report = {"title": "", "paragraphCount": 0, "tableCount": 0, "imageCount": 0,
                       "warnings": [], "errors": [], "resources": []}
        self.root = None
        self.styles, self.images, self.tables = {}, {}, {}
        self.resource_bytes = 0
        self.table_cells = 0
        self.default_font = str(self.config.get("defaultFont") or "Microsoft YaHei")[:200]
        try:
            size = float(self.config.get("defaultFontSizePt", 11))
            self.default_size = min(72, max(6, size)) if math.isfinite(size) else 11
        except (TypeError, ValueError, OverflowError):
            self.default_size = 11
        self.defaults = {"font-family": self.default_font, "font-size": f"{self.default_size}pt",
                         "font-weight": "normal", "font-style": "normal", "color": "black",
                         "line-height": "normal", "text-align": "left", "white-space": "normal",
                         "text-decoration": "none", "visibility": "visible"}
        self._parse(source, resources)

    def note(self, kind, message, node=None):
        message = f"{_location(node)}: {message}" if node is not None else message
        notes = self.report[kind]
        if message not in notes and len(notes) < 200:
            notes.append(message)

    def _parse(self, source, resources):
        if not isinstance(source, bytes) or not source:
            self.note("errors", "HTML source must be nonempty UTF-8 bytes.")
            return
        if len(source) > _limit(self.config, "maxSourceBytes"):
            self.note("errors", "HTML exceeds maxSourceBytes.")
            return
        try:
            text = source.decode("utf-8-sig", errors="strict")
        except UnicodeDecodeError:
            self.note("errors", "HTML must use UTF-8 encoding.")
            return
        safe_text = re.sub(r"<!doctype\s+html\s*>", "", text, flags=re.I)
        if re.search(r"<!\s*(?:DOCTYPE|ENTITY)\b|<\?", safe_text, re.I):
            self.note("errors", "DTD, entity declarations and processing instructions are forbidden.")
            return
        if re.search(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", text):
            self.note("errors", "HTML contains invalid control characters.")
            return
        try:
            parser = html.HTMLParser(no_network=True, recover=True, remove_comments=True)
            self.root = html.document_fromstring(text, parser=parser)
            # libxml's HTML4 parser calls valid HTML5/SVG tags "unknown".
            # Other recoveries can discard content, so do not export them.
            for issue in parser.error_log:
                if issue.type_name != "HTML_UNKNOWN_TAG":
                    self.note("errors", f"HTML structure cannot be parsed reliably (line {issue.line}): {issue.message}")
            if self.report["errors"]:
                return
        except (etree.LxmlError, ValueError) as exc:
            self.note("errors", f"HTML structure cannot be parsed reliably: {exc}")
            return
        # libxml can place a leading SVG fragment after <style> inside <head>.
        # Move body content out explicitly instead of silently dropping it.
        head = self.root.find("head")
        if head is not None:
            misplaced = [child for child in head if child.tag not in IGNORED | {"base", "link", "script"}]
            if misplaced:
                body = self.root.find("body")
                if body is None:
                    body = etree.SubElement(self.root, "body")
                for index, child in enumerate(misplaced):
                    body.insert(index, child)
        nodes = list(self.root.iter())
        if len(nodes) > _limit(self.config, "maxElements"):
            self.note("errors", "HTML exceeds maxElements.")
            return
        depths = {}
        for node in nodes:
            depths[node] = depths.get(node.getparent(), 0) + 1
            if depths[node] > MAX_DEPTH:
                self.note("errors", f"Structure exceeds maximum depth {MAX_DEPTH}.", node)
                return
        title = self.root.find(".//title")
        if title is not None:
            self.report["title"] = SPACE.sub(" ", "".join(title.itertext())).strip()
        self._cascade(nodes)
        if resources is not None and not isinstance(resources, dict):
            self.note("errors", "Resources must map raw image src values to bytes.")
            return
        if resources is not None:
            if any(not isinstance(value, bytes) for value in resources.values()):
                self.note("errors", "Resource values must be bytes.")
                return
            self.resource_bytes = sum(map(len, resources.values()))
            if self.resource_bytes > _limit(self.config, "maxResourceBytes"):
                self.note("errors", "Images exceed total maxResourceBytes.")
                return
        svg_children = set()
        for node in nodes:
            if node in svg_children:
                continue
            tag = node.tag
            if any(key.lower().startswith("on") for key in node.attrib):
                self.note("errors", "Event handlers are dynamic content; export static HTML/SVG/PNG first.", node)
            if tag in DYNAMIC:
                self.note("errors", f"Dynamic or unsupported <{tag}> content; export static SVG/PNG first.", node)
            if tag == "link" and "stylesheet" in node.get("rel", "").lower():
                self.note("errors", "External stylesheets are blocked; use an internal <style>.", node)
            if tag == "base":
                self.note("errors", "Base URLs are unsupported; use explicit workspace-relative image src values.", node)
            if tag == "svg":
                svg_children.update(node.iterdescendants())
            if self.hidden(node):
                continue
            if tag == "svg":
                self.report["imageCount"] += 1
                self._svg(node)
            elif tag == "img":
                self.report["imageCount"] += 1
                self._image(node, resources)
            elif tag == "table":
                self.report["tableCount"] += 1
                self._table(node)
            elif tag in {"ul", "ol"}:
                self._list(node)
            elif tag in {"td", "th"} and (node.getparent() is None or node.getparent().tag != "tr"):
                self.note("errors", "Table cell is outside a table row.", node)
            elif tag == "tr" and (node.getparent() is None or node.getparent().tag not in {"table", "thead", "tbody", "tfoot"}):
                self.note("errors", "Table row is outside a table.", node)
            elif tag == "a":
                href = node.get("href", "").strip()
                try:
                    scheme = urlsplit(href).scheme.lower()
                except ValueError:
                    scheme = "invalid"
                if href and not href.startswith("#") and scheme not in {"http", "https", "mailto", "tel"}:
                    self.note("errors", "Hyperlink target must be http(s), mailto, tel or a document anchor.", node)
                if any(ord(char) < 32 for char in href):
                    self.note("errors", "Hyperlink target contains control characters.", node)
            elif tag not in BLOCKS | PARAGRAPHS | INLINE | IGNORED | DYNAMIC | {"li", "tr", "td", "th", "thead", "tbody", "tfoot", "caption", "hr", "link", "base"}:
                self.note("warnings", f"Unsupported <{tag}> semantics; content is retained in document order.", node)
        if not self.report["title"]:
            heading = self.root.find(".//h1")
            if heading is not None:
                self.report["title"] = SPACE.sub(" ", "".join(heading.itertext())).strip()

    def hidden(self, node):
        return self.styles.get(node, {}).get("_hidden", False)

    def _cascade(self, nodes):
        matcher = cssselect2.Matcher()
        def stylesheet(content, owner):
            for rule in tinycss2.parse_stylesheet(content, skip_comments=True, skip_whitespace=True):
                if rule.type == "at-rule":
                    media = tinycss2.serialize(rule.prelude).strip().lower()
                    if rule.lower_at_keyword == "media" and media in {"print", "all"} and rule.content is not None:
                        stylesheet(rule.content, owner)
                    else:
                        kind = "errors" if rule.lower_at_keyword in {"import", "font-face"} else "warnings"
                        self.note(kind, f"Unsupported CSS @{rule.lower_at_keyword}; keep styles internal and static.", owner)
                elif rule.type == "qualified-rule":
                    selector_text = tinycss2.serialize(rule.prelude).strip()
                    if len(selector_text) > 4096:
                        self.note("errors", "CSS selector is too complex.", owner)
                        continue
                    try:
                        selectors = cssselect2.compile_selector_list(rule.prelude)
                        for selector in selectors:
                            if selector.pseudo_element:
                                self.note("warnings", f"CSS generated content/decorations ({selector_text}) are unsupported.", owner)
                            else:
                                matcher.add_selector(selector, _declarations(rule.content))
                    except cssselect2.SelectorError:
                        self.note("warnings", f"Unsupported CSS selector: {selector_text[:160]}", owner)
                elif rule.type == "error":
                    self.note("warnings", "Invalid CSS rule was ignored.", owner)
        for node in nodes:
            if node.tag == "style":
                stylesheet(node.text or "", node)
        wrapper = cssselect2.ElementWrapper.from_html_root(self.root)
        for wrapped in wrapper.iter_subtree():
            node = wrapped.etree_element
            parent = self.styles.get(node.getparent(), self.defaults)
            style = {key: parent.get(key, self.defaults.get(key, "")) for key in INHERITED}
            tag = node.tag
            in_svg = tag == "svg" or any(ancestor.tag == "svg" for ancestor in node.iterancestors())
            if in_svg:
                for name in ("font-family", "font-size", "font-weight", "font-style", "color"):
                    if node.get(name):
                        style[name] = node.get(name)
            if tag in {"b", "strong", "th"}:
                style["font-weight"] = "bold"
            if tag in {"i", "em", "cite", "var", "address"}:
                style["font-style"] = "italic"
            if tag in {"u", "ins"}:
                style["text-decoration"] = "underline"
            if tag in {"s", "strike", "del"}:
                style["text-decoration"] = "line-through"
            if tag in {"sup", "sub"}:
                style.update({"vertical-align": "super" if tag == "sup" else "sub", "font-size": ".83em"})
            if tag in {"small", "big"}:
                style["font-size"] = ".83em" if tag == "small" else "1.2em"
            if tag in {"pre", "code", "kbd", "samp"}:
                style["font-family"] = "Consolas"
            if tag == "pre":
                style["white-space"] = "pre"
            if tag == "mark":
                style["background-color"] = "yellow"
            if tag in {"p", "pre"}:
                style.update({"margin-top": "0pt", "margin-bottom": "8pt"})
            if tag in {"h1", "h2", "h3", "h4", "h5", "h6"}:
                style.update({"font-size": f"{(2, 1.5, 1.17, 1, .83, .67)[int(tag[1]) - 1]}em",
                              "font-weight": "bold", "margin-top": "12pt", "margin-bottom": "6pt"})
            if tag == "blockquote":
                style.update({"margin-left": "24pt", "margin-right": "12pt"})
            if tag == "th":
                style["text-align"] = "center"
            if tag in {"ol", "ul"}:
                style["list-style-type"] = {"1": "decimal", "a": "lower-alpha", "A": "upper-alpha", "i": "lower-roman", "I": "upper-roman"}.get(node.get("type"), "decimal" if tag == "ol" else "disc")
            if node.get("align"):
                style["text-align"] = node.get("align")
            if node.get("bgcolor"):
                style["background-color"] = node.get("bgcolor")
            if node.get("valign"):
                style["vertical-align"] = node.get("valign")
            winners = {}
            for specificity, order, pseudo, declarations in matcher.match(wrapped):
                for index, (name, value, important) in enumerate(declarations):
                    rank = (important, 0, specificity, order, index)
                    if name not in winners or rank >= winners[name][0]:
                        winners[name] = (rank, value)
            for index, (name, value, important) in enumerate(_declarations(node.get("style", ""))):
                rank = (important, 1, (0, 0, 0), 0, index)
                if name not in winners or rank >= winners[name][0]:
                    winners[name] = (rank, value)
            for name, (_, value) in winners.items():
                if value == "inherit" or (value == "unset" and name in INHERITED):
                    value = parent.get(name, self.defaults.get(name, ""))
                elif value in {"initial", "unset", "revert"}:
                    value = self.defaults.get(name, "")
                style[name] = value
            parent_size = _length(parent.get("font-size", "11pt")) or self.default_size
            size_value = style.get("font-size", "11pt")
            keywords = {"xx-small": 7, "x-small": 8, "small": 9, "medium": self.default_size, "large": 14,
                        "x-large": 18, "xx-large": 24, "smaller": parent_size / 1.2, "larger": parent_size * 1.2}
            size = keywords.get(size_value, _length(size_value, parent_size, parent_size))
            if size is None or size <= 0 or size > 400:
                self.note("warnings", f"Unsupported font-size {size_value}; using default.", node)
                size = self.default_size
            style["font-size"] = f"{size}pt"
            if style.get("font-weight") in {"bolder", "lighter"}:
                style["font-weight"] = "bold" if style["font-weight"] == "bolder" else "normal"
            line = style.get("line-height", "normal")
            if line != "normal" and not re.fullmatch(r"\d*\.?\d+", line):
                length = _length(line, size, size)
                if length is not None and length > 0:
                    style["line-height"] = f"{min(1000, length)}pt"
            style["_hidden"] = parent.get("_hidden", False) or style.get("display") == "none" or "hidden" in node.attrib
            style["_baseline"] = style.get("vertical-align", parent.get("_baseline"))
            style["_declared"] = {name: style.get(name, value) for name, (_, value) in winners.items()}
            self.styles[node] = style
            if in_svg:
                continue
            for name, (_, value) in winners.items():
                if name not in SUPPORTED:
                    self.note("warnings", f"Unsupported CSS decoration/property {name}: {value[:120]}.", node)
                elif ((name == "display" and value in {"flex", "inline-flex", "grid", "inline-grid"}) or
                      (name == "float" and value != "none") or (name == "position" and value in {"absolute", "fixed", "sticky", "relative"})):
                    self.note("warnings", f"CSS layout simplified ({name}: {value}); content follows document order.", node)
                elif name == "visibility" and value != "visible":
                    self.note("warnings", "CSS visibility is simplified; use display:none to omit content.", node)
                elif name in {"width", "height", "max-width", "max-height"} and tag not in {"img", "svg", "table", "td", "th", "col", "colgroup"}:
                    self.note("warnings", f"CSS layout simplified ({name}); text flows within the page.", node)
                elif name.startswith("padding-") and tag not in {"td", "th"}:
                    self.note("warnings", f"CSS layout simplified ({name}); only table cell padding is supported.", node)
                elif name in {"border-spacing", "table-layout"} and value not in {"0", "0px", "fixed"}:
                    self.note("warnings", f"CSS layout simplified ({name}); Word uses a fixed table grid.", node)
                elif name == "text-decoration" and any(part not in {"none", "underline", "line-through"} for part in value.split()):
                    self.note("warnings", f"Unsupported text decoration ({value}); only underline/line-through are retained.", node)
                elif name == "text-align" and value not in {"left", "right", "center", "justify", "start", "end"}:
                    self.note("warnings", f"Unsupported text-align ({value}); using left alignment.", node)
                elif name == "white-space" and value not in {"normal", "nowrap", "pre", "pre-wrap", "pre-line", "break-spaces"}:
                    self.note("warnings", f"Unsupported white-space ({value}); collapsing ordinary whitespace.", node)
                elif name == "white-space" and value == "nowrap":
                    self.note("warnings", "CSS nowrap is simplified; Word may wrap text at the page edge.", node)
                elif name in {"break-before", "break-after", "page-break-before", "page-break-after"} and value in {"left", "right", "recto", "verso"}:
                    self.note("warnings", f"CSS {name}:{value} becomes a page break without page-side selection.", node)
                elif name == "vertical-align" and value not in {"baseline", "super", "sub", "top", "middle", "bottom"}:
                    self.note("warnings", f"Unsupported vertical-align ({value}); using the normal baseline.", node)
                elif name in {"height", "max-height"} and tag in {"table", "td", "th", "col", "colgroup"}:
                    self.note("warnings", f"CSS layout simplified ({name}); table height follows editable content.", node)
                elif name.startswith("border-") and name.endswith(("-style", "-width", "-color")) and tag in BLOCKS:
                    self.note("warnings", "Container border decoration is simplified; put borders on paragraphs or table cells.", node)
                elif name.startswith("margin-") and tag in INLINE:
                    self.note("warnings", "Inline margin decoration is unsupported; use paragraph indentation/spacing.", node)
                elif (name in {"width", "height", "max-width", "max-height", "text-indent"} or name.startswith(("margin-", "padding-"))) and value not in {"", "auto", "none"} and _length(value) is None:
                    self.note("warnings", f"Unsupported CSS length {name}: {value[:120]}.", node)
                elif (name in {"color", "background-color"} or name.endswith("-color")) and value not in {"transparent", "currentcolor", "currentColor"} and _color(value) is None:
                    self.note("warnings", f"Unsupported CSS color {name}: {value[:120]}.", node)

    def _resource_size(self, size, node):
        self.resource_bytes += size
        if self.resource_bytes > _limit(self.config, "maxResourceBytes"):
            self.note("errors", "Images exceed total maxResourceBytes.", node)
            return False
        return True

    def _raster(self, data, node):
        try:
            with Image.open(BytesIO(data)) as picture:
                if picture.format not in {"PNG", "JPEG"}:
                    raise ValueError("only PNG and JPEG images are supported")
                width, height = picture.size
                if width * height > _limit(self.config, "maxImagePixels"):
                    raise ValueError("image exceeds maxImagePixels")
                if width < 1 or height < 1:
                    raise ValueError("empty image")
                if getattr(picture, "n_frames", 1) > 1:
                    raise ValueError("animated images must be exported as static PNG")
                picture.verify()
            # verify() does not decode JPEG data. Decode before accepting it.
            with Image.open(BytesIO(data)) as picture:
                picture.load()
            return data, width * .75, height * .75
        except (OSError, ValueError, SyntaxError, Image.DecompressionBombError) as exc:
            self.note("errors", f"Invalid image: {exc}", node)
            return None

    def _image(self, node, resources):
        raw = node.get("src", "")
        src = raw.strip()
        if not src:
            self.note("errors", "Image src is missing.", node)
            return
        if node.get("srcset"):
            self.note("warnings", "Responsive srcset is ignored; using src.", node)
        if src.lower().startswith("data:"):
            match = re.fullmatch(r"data:image/(png|jpeg);base64,([A-Za-z0-9+/=\s]+)", src, flags=re.I)
            if not match:
                self.note("errors", "Only base64 PNG/JPEG data images are supported.", node)
                return
            try:
                data = base64.b64decode(SPACE.sub("", match[2]), validate=True)
            except (ValueError, binascii.Error):
                self.note("errors", "Image data is invalid base64.", node)
                return
            if not self._resource_size(len(data), node):
                return
        else:
            decoded = unquote(src).replace("\\", "/")
            try:
                parsed = urlsplit(decoded)
                external = bool(parsed.scheme or parsed.netloc or parsed.query or parsed.fragment)
            except ValueError:
                external = True
            if external or decoded.startswith(("/", "#")) or any(ord(char) < 32 for char in decoded):
                self.note("errors", "External or unsafe image src is blocked; use a workspace-relative path.", node)
                return
            # Parent-relative paths are valid when the Host API resolves them
            # inside the workspace. This converter receives only resource bytes.
            if raw not in self.report["resources"]:
                self.report["resources"].append(raw)
            if resources is None:
                return
            if raw not in resources:
                self.note("errors", f"Missing image resource: {raw}", node)
                return
            data = resources[raw]
        image = self._raster(data, node)
        if image:
            self.images[node] = image

    def _svg(self, node):
        """Reject all non-fragment references, including CSS-escaped URLs."""
        import resvg_py

        invalid = False
        references = []
        def reject(message, element=node):
            nonlocal invalid
            invalid = True
            self.note("errors", message, element)

        def tokens_safe(tokens, owner):
            for token in tokens:
                if token.type == "url":
                    if not token.value.strip().startswith("#"):
                        reject("SVG external references are blocked.", owner)
                    else:
                        references.append((owner, unquote(token.value.strip()[1:])))
                if token.type == "function":
                    if token.lower_name == "url":
                        arguments = [part for part in token.arguments if part.type != "whitespace"]
                        value = arguments[0].value if len(arguments) == 1 and arguments[0].type in {"string", "ident", "hash"} else ""
                        if not str(value).startswith("#"):
                            reject("SVG external references are blocked.", owner)
                        else:
                            references.append((owner, unquote(value[1:])))
                    tokens_safe(token.arguments, owner)
                elif token.type == "at-keyword" and token.lower_value in {"import", "font-face"}:
                    reject("SVG CSS imports and external fonts are blocked.", owner)
                elif token.type in {"{} block", "[] block", "() block"}:
                    tokens_safe(token.content, owner)
                elif token.type == "error":
                    reject("SVG contains invalid CSS reference syntax.", owner)

        for child in node.iter():
            tag = child.tag.lower()
            if tag in {"script", "foreignobject", "animate", "animatemotion", "animatetransform", "set", "discard", "iframe", "audio", "video"}:
                reject(f"SVG dynamic/unsupported <{tag}>; export a static chart first.", child)
            for key, value in child.attrib.items():
                if key.lower().startswith("on") or key.lower() in {"xml:base", "base"}:
                    reject("SVG event handlers and base paths are blocked.", child)
                if key.lower().split(":")[-1] in {"href", "src"}:
                    if not value.strip().startswith("#"):
                        reject("SVG external references are blocked.", child)
                    else:
                        references.append((child, unquote(value.strip()[1:])))
                tokens_safe(tinycss2.parse_component_value_list(value), child)
            if tag == "style":
                tokens_safe(tinycss2.parse_component_value_list(child.text or ""), child)
            for value in self.styles[child].get("_declared", {}).values():
                tokens_safe(tinycss2.parse_component_value_list(value), child)
        ids = {}
        for child in node.iter():
            ident = child.get("id")
            if ident:
                if ident in ids:
                    reject("SVG has duplicate IDs; references are ambiguous.", child)
                ids[ident] = child
        edges = {ident: set() for ident in ids}
        for owner, target in references:
            if target not in ids:
                reject(f"SVG reference #{target} is missing.", owner)
                continue
            for ancestor in [owner, *owner.iterancestors()]:
                if ancestor.get("id") in edges:
                    edges[ancestor.get("id")].add(target)
                if ancestor is node:
                    break
        visiting, visited = set(), set()
        def visit(ident, depth=0):
            if ident in visiting or depth > MAX_DEPTH:
                reject("SVG cyclic or excessively deep references cannot be rendered reliably.", ids[ident])
                return
            if ident in visited:
                return
            visiting.add(ident)
            for target in edges[ident]:
                visit(target, depth + 1)
            visiting.remove(ident)
            visited.add(ident)
        for ident in ids:
            visit(ident)
        if invalid:
            return
        # HTML parsing folds SVG camelCase names. Restore the SVG vocabulary.
        tags = "altGlyph altGlyphDef altGlyphItem animateColor animateMotion animateTransform clipPath feBlend feColorMatrix feComponentTransfer feComposite feConvolveMatrix feDiffuseLighting feDisplacementMap feDistantLight feDropShadow feFlood feFuncA feFuncB feFuncG feFuncR feGaussianBlur feImage feMerge feMergeNode feMorphology feOffset fePointLight feSpecularLighting feSpotLight feTile feTurbulence foreignObject glyphRef linearGradient radialGradient textPath".split()
        attrs = "attributeName attributeType baseFrequency baseProfile calcMode clipPathUnits diffuseConstant edgeMode filterUnits glyphRef gradientTransform gradientUnits kernelMatrix kernelUnitLength keyPoints keySplines keyTimes lengthAdjust limitingConeAngle markerHeight markerUnits markerWidth maskContentUnits maskUnits numOctaves pathLength patternContentUnits patternTransform patternUnits pointsAtX pointsAtY pointsAtZ preserveAlpha preserveAspectRatio primitiveUnits refX refY repeatCount repeatDur requiredExtensions requiredFeatures specularConstant specularExponent spreadMethod startOffset stdDeviation stitchTiles surfaceScale systemLanguage tableValues targetX targetY textLength viewBox xChannelSelector yChannelSelector zoomAndPan".split()
        tag_names, attr_names = {x.lower(): x for x in tags}, {x.lower(): x for x in attrs}
        svg_ns, xlink_ns = "http://www.w3.org/2000/svg", "http://www.w3.org/1999/xlink"
        def copy(original):
            result = etree.Element(f"{{{svg_ns}}}{tag_names.get(original.tag, original.tag)}", nsmap={None: svg_ns, "xlink": xlink_ns} if original is node else None)
            for name, value in original.attrib.items():
                if name == "xmlns" or name.startswith("xmlns:"):
                    continue
                if name.startswith("xlink:"):
                    name = f"{{{xlink_ns}}}{name[6:]}"
                elif name.startswith("xml:"):
                    name = f"{{http://www.w3.org/XML/1998/namespace}}{name[4:]}"
                else:
                    name = attr_names.get(name, name)
                result.set(name, value)
            declared = self.styles[original].get("_declared", {})
            if declared:
                result.set("style", ";".join(f"{key}:{value} !important" for key, value in declared.items()))
            result.text, result.tail = original.text, original.tail if original is not node else None
            for child in original:
                result.append(copy(child))
            return result
        try:
            root = copy(node)
            style = self.styles[node]
            viewbox = [float(part) for part in re.split(r"[\s,]+", node.get("viewbox", "").strip()) if part]
            if viewbox and (len(viewbox) != 4 or not all(math.isfinite(n) for n in viewbox) or min(viewbox[2:]) <= 0):
                raise ValueError("invalid SVG viewBox")
            width = _length(style.get("width", node.get("width", "")))
            height = _length(style.get("height", node.get("height", "")), reference=PAGE_HEIGHT)
            if width is None and height is None:
                width, height = (viewbox[2] * .75, viewbox[3] * .75) if viewbox else (225, 112.5)
            elif width is None:
                width = height * (viewbox[2] / viewbox[3] if viewbox else 2)
            elif height is None:
                height = width * (viewbox[3] / viewbox[2] if viewbox else .5)
            pixels = (max(1, math.ceil(width / .75)), max(1, math.ceil(height / .75)))
            if min(width, height) <= 0 or pixels[0] * pixels[1] > _limit(self.config, "maxImagePixels"):
                raise ValueError("SVG exceeds maxImagePixels or has invalid dimensions")
            root.set("width", str(pixels[0]))
            root.set("height", str(pixels[1]))
            # Parent HTML styles do not exist in the standalone SVG document.
            for key in ("font-family", "font-size", "font-weight", "font-style", "color"):
                if key not in root.attrib:
                    root.set(key, style[key])
            data = resvg_py.svg_to_bytes(svg_string=etree.tostring(root, encoding="unicode"),
                                         width=pixels[0], height=pixels[1], log_information=False,
                                         font_family=self.default_font)
            if self._resource_size(len(data), node):
                self.images[node] = (data, width, height)
        except (ValueError, OSError, etree.LxmlError, OverflowError) as exc:
            self.note("errors", f"Invalid SVG chart: {exc}", node)

    def _list(self, node):
        depth = sum(parent.tag in {"ol", "ul"} for parent in node.iterancestors())
        if depth > 8:
            self.note("errors", "Word supports at most 9 native list levels.", node)
        if node.get("reversed") is not None:
            self.note("errors", "Reversed lists are unsupported; use an ordinary ordered list.", node)
        if self.styles[node].get("list-style-type") not in LIST_FORMATS:
            self.note("warnings", "Unsupported list-style-type; using decimal or disc.", node)
        if (node.text or "").strip() or any(child.tag != "li" or (child.tail or "").strip() for child in node):
            self.note("errors", "List structure must contain only <li> elements.", node)
        for child in [node, *node]:
            raw = child.get("start" if child is node else "value")
            if raw is not None:
                try:
                    if abs(int(raw)) > 1_000_000:
                        raise ValueError()
                except ValueError:
                    self.note("errors", "List start/value must be an integer within ±1000000.", child)

    def _table(self, node):
        rows, group_ends = [], {}
        for child in node:
            if self.hidden(child):
                continue
            if child.tag == "tr":
                rows.append(child)
            elif child.tag in {"thead", "tbody", "tfoot"}:
                if (child.text or "").strip() or any(row.tag != "tr" or (row.tail or "").strip() for row in child):
                    self.note("errors", "Table row group contains non-row content.", child)
                    return
                grouped = [row for row in child if not self.hidden(row)]
                rows.extend(grouped)
                group_ends.update({row: len(rows) for row in grouped})
            elif child.tag not in {"caption", "colgroup", "col", "style"}:
                self.note("errors", "Table contains content outside rows/caption.", child)
                return
            if (child.tail or "").strip():
                self.note("errors", "Table contains text outside cells.", child)
                return
        if (node.text or "").strip() or not rows:
            self.note("errors", "Table needs rows and text inside cells.", node)
            return
        occupied, cells, columns = set(), [], 0
        limit = _limit(self.config, "maxTableCells")
        for row_index, row in enumerate(rows):
            if (row.text or "").strip():
                self.note("errors", "Table row contains text outside cells.", row)
                return
            column = 0
            for cell in row:
                if cell.tag not in {"td", "th"} or (cell.tail or "").strip():
                    self.note("errors", "Table row must contain only cells.", row)
                    return
                if self.hidden(cell):
                    continue
                while (row_index, column) in occupied:
                    column += 1
                try:
                    colspan, rowspan = int(cell.get("colspan", 1)), int(cell.get("rowspan", 1))
                    end = group_ends.get(row, len(rows))
                    if rowspan == 0:
                        rowspan = end - row_index
                    if colspan < 1 or rowspan < 1 or column + colspan > 63 or row_index + rowspan > end or colspan * rowspan > limit:
                        raise ValueError()
                except ValueError:
                    self.note("errors", "Invalid/oversized table span or more than 63 Word columns (maxTableCells).", cell)
                    return
                for y in range(row_index, row_index + rowspan):
                    for x in range(column, column + colspan):
                        if (y, x) in occupied:
                            self.note("errors", "Overlapping table spans cannot be converted reliably.", cell)
                            return
                        occupied.add((y, x))
                cells.append((cell, row_index, column, rowspan, colspan))
                column += colspan
                columns = max(columns, column)
                if len(occupied) + self.table_cells > limit:
                    self.note("errors", "Tables exceed maxTableCells.", node)
                    return
        if not columns or len(rows) * columns != len(occupied):
            self.note("errors", "Ragged table has missing cells; use explicit colspan/rowspan.", node)
            return
        self.table_cells += len(occupied)
        self.tables[node] = (rows, columns, cells)


class _Flow:
    """A block's current paragraph and whitespace state, shared by inline children."""
    def __init__(self, renderer, container, width, style, indent=0, right=0, numbering=None):
        self.renderer, self.container, self.width, self.style = renderer, container, width, style
        self.indent, self.right, self.numbering = indent, right, numbering
        self.paragraph = None
        self.pending_space = False
        self.line_start = True
        self.first = True
        self.paragraphs = []
        self.pending_before = 0

    def end(self):
        self.paragraph = None
        self.pending_space = False
        self.line_start = True

    def ensure(self, style=None, tag=None):
        if self.paragraph is not None:
            if style is not None and not self.paragraph.text and not self.paragraph._p.xpath(".//w:drawing|.//w:hyperlink"):
                self.renderer.paragraph(self.paragraph, style, tag, self.indent, self.right, self.width)
            return self.paragraph
        style = style or self.style
        # Every Word cell starts with a mandatory empty paragraph.
        paragraphs = self.container.paragraphs
        if hasattr(self.container, "_tc") and len(paragraphs) == 1 and not len(paragraphs[0]._p) and len(self.container._tc) <= 2:
            self.paragraph = paragraphs[0]
        else:
            self.paragraph = self.container.add_paragraph()
        self.paragraphs.append(self.paragraph)
        self.renderer.source.report["paragraphCount"] += 1
        self.renderer.paragraph(self.paragraph, style, tag, self.indent, self.right, self.width)
        if self.pending_before:
            current = self.paragraph.paragraph_format.space_before
            self.paragraph.paragraph_format.space_before = Pt(max(current.pt if current else 0, self.pending_before))
            self.pending_before = 0
        if self.numbering:
            num_id, level = self.numbering
            if self.first:
                num_pr = self.paragraph._p.get_or_add_pPr().get_or_add_numPr()
                num_pr.get_or_add_numId().val = num_id
                num_pr.get_or_add_ilvl().val = level
                self.paragraph.paragraph_format.left_indent = Pt(self.indent + 18 * (level + 1))
                self.paragraph.paragraph_format.first_line_indent = Pt(-9)
            else:
                self.paragraph.paragraph_format.left_indent = Pt(self.indent + 18 * (level + 1))
        self.first = False
        return self.paragraph

    def text(self, text, style, link=None):
        if not text:
            return
        whitespace = style.get("white-space", "normal")
        if whitespace in {"pre", "pre-wrap", "break-spaces"}:
            value = text.replace("\r\n", "\n").replace("\r", "\n")
            if self.pending_space and not self.line_start:
                value = " " + value
            self.pending_space = False
        elif whitespace == "pre-line":
            parts = text.replace("\r\n", "\n").replace("\r", "\n").split("\n")
            for index, part in enumerate(parts):
                self.text(part, {**style, "white-space": "normal"}, link)
                if index < len(parts) - 1:
                    self.br(style)
            return
        else:
            collapsed = SPACE.sub(" ", text)
            value = collapsed.strip(" ")
            leading = self.pending_space or collapsed.startswith(" ")
            self.pending_space = collapsed.endswith(" ")
            if not value:
                self.pending_space = leading or self.pending_space
                return
            if leading and not self.line_start:
                value = " " + value
        if value:
            paragraph = self.ensure()
            self.renderer.run(paragraph, value, style, link)
            self.line_start = value.endswith("\n")

    def br(self, style):
        self.renderer.run(self.ensure(), "\n", style)
        self.pending_space, self.line_start = False, True

    def picture(self, node, style):
        data = self.renderer.source.images.get(node)
        if data is None:
            self.ensure()
            return  # Discovery inspection has no host resource bytes yet.
        if self.pending_space and not self.line_start:
            self.renderer.run(self.ensure(), " ", style)
        content, original_width, original_height = data
        font = _length(style["font-size"]) or 11
        available = max(12, self.width - self.indent - self.right - (18 * (self.numbering[1] + 1) if self.numbering else 0))
        requested_width = _length(style.get("width", node.get("width", "")), font, available)
        requested_height = _length(style.get("height", node.get("height", "")), font, PAGE_HEIGHT)
        ratio = requested_width / original_width if requested_width and requested_width > 0 else 1
        if requested_height and requested_height > 0:
            ratio = min(ratio, requested_height / original_height) if requested_width else requested_height / original_height
        max_width = _length(style.get("max-width", ""), font, available)
        max_height = _length(style.get("max-height", ""), font, PAGE_HEIGHT)
        ratio = min(ratio, available / original_width, (PAGE_HEIGHT - 12) / original_height)
        if max_width and max_width > 0:
            ratio = min(ratio, max_width / original_width)
        if max_height and max_height > 0:
            ratio = min(ratio, max_height / original_height)
        paragraph = self.ensure()
        if paragraph.paragraph_format.line_spacing_rule == WD_LINE_SPACING.EXACTLY:
            paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.AT_LEAST
        run = paragraph.add_run()
        picture = run.add_picture(BytesIO(content), width=Pt(max(.1, original_width * ratio)), height=Pt(max(.1, original_height * ratio)))
        picture._inline.docPr.set("descr", node.get("alt", node.get("aria-label", "")))
        self.pending_space, self.line_start = False, False


class _Renderer:
    def __init__(self, source):
        self.source = source
        self.document = Document()
        section = self.document.sections[0]
        section.page_width, section.page_height = Mm(210), Mm(297)
        section.top_margin = section.bottom_margin = section.left_margin = section.right_margin = Mm(20)
        for grid in section._sectPr.findall(qn("w:docGrid")):
            section._sectPr.remove(grid)
        normal = self.document.styles["Normal"]
        normal.font.name, normal.font.size = source.default_font, Pt(source.default_size)
        normal.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), source.default_font)
        normal.paragraph_format.space_after = Pt(0)
        normal_properties = normal.element.get_or_add_pPr()
        grid_setting = _element(normal_properties, "w:snapToGrid", val="0")
        normal_properties.insert_element_before(grid_setting, "w:spacing", "w:ind", "w:jc", "w:rPr")
        self.document.core_properties.title = source.report["title"][:255]
        self.bookmarks = {}
        self.bookmarks_written = set()
        for node in source.root.iter():
            ident = node.get("id") or (node.get("name") if node.tag == "a" else None)
            if ident and ident not in self.bookmarks:
                self.bookmarks[ident] = f"html_{len(self.bookmarks) + 1}"
        self.next_num = max([0, *(int(node.get(qn("w:numId"))) for node in self.document.part.numbering_part.element.findall(qn("w:num")))]) + 1
        self.next_abstract = max([0, *(int(node.get(qn("w:abstractNumId"))) for node in self.document.part.numbering_part.element.findall(qn("w:abstractNum")))]) + 1

    def render(self):
        self.source.report["paragraphCount"] = 0
        body = self.source.root.find("body")
        body = body if body is not None else self.source.root
        flow = _Flow(self, self.document, PAGE_WIDTH, self.source.styles[body])
        self.walk(body, flow)
        return self.document

    def paragraph(self, paragraph, style, tag, indent, right, width):
        if tag and re.fullmatch(r"h[1-6]", tag):
            paragraph.style = f"Heading {tag[1]}"
            paragraph.paragraph_format.keep_with_next = True
        formatting = paragraph.paragraph_format
        formatting.widow_control = True
        alignment = {"left": WD_ALIGN_PARAGRAPH.LEFT, "start": WD_ALIGN_PARAGRAPH.LEFT,
                     "center": WD_ALIGN_PARAGRAPH.CENTER, "right": WD_ALIGN_PARAGRAPH.RIGHT,
                     "end": WD_ALIGN_PARAGRAPH.RIGHT, "justify": WD_ALIGN_PARAGRAPH.JUSTIFY}
        formatting.alignment = alignment.get(style.get("text-align"), WD_ALIGN_PARAGRAPH.LEFT)
        size = _length(style.get("font-size", "11pt")) or 11
        for property_name, attribute in (("margin-top", "space_before"), ("margin-bottom", "space_after")):
            value = _length(style.get(property_name, "0"), size, width)
            if value is not None:
                setattr(formatting, attribute, Pt(max(0, min(value, PAGE_HEIGHT))))
        left_margin = _length(style.get("margin-left", "0"), size, width) or 0
        right_margin = _length(style.get("margin-right", "0"), size, width) or 0
        formatting.left_indent = Pt(max(0, min(indent + left_margin, width - 12)))
        formatting.right_indent = Pt(max(0, min(right + right_margin, width - 12)))
        first = _length(style.get("text-indent", "0"), size, width) or 0
        formatting.first_line_indent = Pt(max(-width, min(first, width)))
        line = style.get("line-height", "normal")
        if re.fullmatch(r"\d*\.?\d+", line):
            # CSS numbers multiply font size, not Word's font-dependent single
            # spacing. A minimum line box lets larger glyphs and pictures fit.
            formatting.line_spacing = Pt(size * max(.5, min(10, float(line))))
            formatting.line_spacing_rule = WD_LINE_SPACING.AT_LEAST
        else:
            amount = _length(line, size, size)
            if amount and amount > 0:
                formatting.line_spacing = Pt(min(amount, 1000))
                formatting.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        before = style.get("break-before", style.get("page-break-before", ""))
        if before in {"always", "page", "left", "right", "recto", "verso"}:
            formatting.page_break_before = True
        elif before == "avoid":
            formatting.keep_with_next = True
        if style.get("break-inside", style.get("page-break-inside")) in {"avoid", "avoid-page"}:
            formatting.keep_together = True
        fill = _color(style.get("background-color", ""))
        properties = paragraph._p.get_or_add_pPr()
        # The paragraph mark also contributes to Word's minimum line height.
        # Keep it at the CSS paragraph size instead of the template's size.
        for previous in properties.findall(qn("w:rPr")):
            properties.remove(previous)
        mark = _element(properties, "w:rPr")
        family = self.font_family(style)
        _element(mark, "w:rFonts", ascii=family, hAnsi=family, eastAsia=family)
        _element(mark, "w:sz", val=round(size * 2))
        _element(mark, "w:szCs", val=round(size * 2))
        if fill:
            shading = _element(properties, "w:shd", fill=fill, val="clear")
            properties.insert_element_before(shading, "w:tabs", "w:spacing", "w:ind", "w:jc", "w:rPr", "w:sectPr")
        borders = self.borders(style)
        if borders is not None:
            borders.tag = qn("w:pBdr")
            properties.insert_element_before(borders, "w:shd", "w:tabs", "w:spacing", "w:ind", "w:jc", "w:rPr", "w:sectPr")

    def font_family(self, style):
        family = style.get("font-family", self.source.default_font).split(",", 1)[0].strip(" \"'") or self.source.default_font
        return {"sans-serif": self.source.default_font, "serif": "Times New Roman", "monospace": "Consolas"}.get(family, family)[:200]

    def run(self, paragraph, text, style, link=None):
        run = paragraph.add_run(text)
        font = run.font
        font.name = self.font_family(style)
        run._r.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), font.name)
        font.size = Pt(_length(style["font-size"]) or self.source.default_size)
        formatting = paragraph.paragraph_format
        line = style.get("line-height", "normal")
        if formatting.line_spacing_rule == WD_LINE_SPACING.AT_LEAST and re.fullmatch(r"\d*\.?\d+", line):
            requested = font.size.pt * max(.5, min(10, float(line)))
            formatting.line_spacing = Pt(max(formatting.line_spacing.pt, requested))
        elif formatting.line_spacing_rule == WD_LINE_SPACING.EXACTLY and font.size.pt > formatting.line_spacing.pt:
            formatting.line_spacing_rule = WD_LINE_SPACING.AT_LEAST
        weight = style.get("font-weight", "normal")
        font.bold = weight == "bold" or (weight.isdigit() and int(weight) >= 600)
        font.italic = style.get("font-style") in {"italic", "oblique"}
        decoration = style.get("text-decoration", "")
        font.underline = "underline" in decoration or bool(link)
        font.strike = "line-through" in decoration
        vertical = style.get("_baseline", style.get("vertical-align"))
        font.superscript = vertical == "super"
        if vertical == "sub":
            font.subscript = True
        color = _color(style.get("color", "black"))
        if color:
            font.color.rgb = RGBColor.from_string(color)
        fill = _color(style.get("background-color", ""))
        if fill:
            properties = run._r.get_or_add_rPr()
            shading = _element(properties, "w:shd", fill=fill, val="clear")
            properties.insert_element_before(shading, "w:vertAlign", "w:rtl", "w:lang")
        if link:
            hyperlink = OxmlElement("w:hyperlink")
            if link.startswith("#"):
                target = self.bookmarks.get(unquote(link[1:]))
                if target:
                    hyperlink.set(qn("w:anchor"), target)
                else:
                    self.source.note("warnings", f"Document anchor {link} does not exist; link text is retained.")
                    return run
            else:
                hyperlink.set(qn("r:id"), paragraph.part.relate_to(link, RT.HYPERLINK, is_external=True))
            hyperlink.append(run._r)
            paragraph._p.append(hyperlink)
        return run

    def bookmark(self, node, flow):
        ident = node.get("id") or (node.get("name") if node.tag == "a" else None)
        if not ident or ident in self.bookmarks_written:
            return
        self.bookmarks_written.add(ident)
        number = len(self.bookmarks_written)
        paragraph = flow.ensure()
        _element(paragraph._p, "w:bookmarkStart", id=number, name=self.bookmarks[ident])
        _element(paragraph._p, "w:bookmarkEnd", id=number)

    def children(self, node, flow, link=None):
        style = flow.style if node.tag in {"td", "th"} else self.source.styles[node]
        self.bookmark(node, flow)
        flow.text(node.text, style, link)
        for child in node:
            self.walk(child, flow, link)
            flow.text(child.tail, style, link)

    def walk(self, node, flow, link=None):
        tag, style = node.tag, self.source.styles[node]
        if self.source.hidden(node) or tag in IGNORED | {"base", "link"}:
            return
        if tag in {"img", "svg"}:
            is_block = style.get("display") == "block"
            if is_block:
                flow.end()
            self.bookmark(node, flow)
            flow.picture(node, style)
            if is_block:
                flow.end()
        elif tag == "br":
            flow.br(style)
        elif tag == "wbr":
            pass
        elif tag == "hr":
            flow.end()
            paragraph = flow.ensure(style)
            properties = paragraph._p.get_or_add_pPr()
            borders = OxmlElement("w:pBdr")
            _element(borders, "w:bottom", val="single", sz=4, color="AAAAAA")
            properties.insert_element_before(borders, "w:shd", "w:spacing", "w:ind", "w:jc")
            flow.end()
        elif tag == "table":
            flow.end()
            self.table(node, flow)
        elif tag in {"ul", "ol"}:
            flow.end()
            self.list(node, flow)
        elif tag in PARAGRAPHS:
            if flow.paragraph is not None and (flow.paragraph.text or flow.paragraph._p.xpath(".//w:drawing|.//w:hyperlink")):
                flow.end()
            paragraph = flow.ensure(style, tag)
            self.children(node, flow, link)
            self.after(style, paragraph)
            flow.end()
        elif tag in BLOCKS or tag == "li" or style.get("display") in {"block", "flex", "grid", "list-item"}:
            flow.end()
            size = _length(style["font-size"]) or 11
            left = _length(style.get("margin-left", "0"), size, flow.width) or 0
            right = _length(style.get("margin-right", "0"), size, flow.width) or 0
            child_flow = _Flow(self, flow.container, flow.width,
                               {**style, "margin-left": "0", "margin-right": "0"},
                               min(flow.width - 12, flow.indent + left), min(flow.width - 12, flow.right + right), flow.numbering)
            child_flow.first = flow.first
            child_flow.pending_before, flow.pending_before = flow.pending_before, 0
            if style.get("break-before", style.get("page-break-before")) in {"page", "always", "left", "right", "recto", "verso"}:
                child_flow.ensure().paragraph_format.page_break_before = True
            self.children(node, child_flow, link)
            if not child_flow.paragraphs and style.get("break-after", style.get("page-break-after")) in {"page", "always", "left", "right", "recto", "verso"}:
                child_flow.ensure()
            if child_flow.paragraphs:
                first, last = child_flow.paragraphs[0], child_flow.paragraphs[-1]
                if style.get("break-before", style.get("page-break-before")) in {"page", "always", "left", "right", "recto", "verso"}:
                    first.paragraph_format.page_break_before = True
                top = _length(style.get("margin-top", "0"), size, flow.width) or 0
                bottom = _length(style.get("margin-bottom", "0"), size, flow.width) or 0
                first.paragraph_format.space_before = Pt(max(first.paragraph_format.space_before.pt if first.paragraph_format.space_before else 0, min(PAGE_HEIGHT, top)))
                last.paragraph_format.space_after = Pt(max(last.paragraph_format.space_after.pt if last.paragraph_format.space_after else 0, min(PAGE_HEIGHT, bottom)))
                self.after(style, last)
                flow.paragraphs.extend(child_flow.paragraphs)
                flow.first = child_flow.first
            flow.pending_before = child_flow.pending_before
        else:
            self.children(node, flow, node.get("href", "").strip() if tag == "a" else link)

    def after(self, style, paragraph):
        value = style.get("break-after", style.get("page-break-after", ""))
        if value in {"always", "page", "left", "right", "recto", "verso"}:
            paragraph.add_run().add_break(WD_BREAK.PAGE)
        elif value in {"avoid", "avoid-page"}:
            paragraph.paragraph_format.keep_with_next = True

    def numbering(self, kind, level, start):
        root = self.document.part.numbering_part.element
        abstract = OxmlElement("w:abstractNum")
        abstract.set(qn("w:abstractNumId"), str(self.next_abstract))
        # Abstract definitions must precede concrete num instances in OOXML.
        first_num = root.find(qn("w:num"))
        root.insert(root.index(first_num) if first_num is not None else len(root), abstract)
        _element(abstract, "w:multiLevelType", val="multilevel")
        for index in range(9):
            lvl = _element(abstract, "w:lvl", ilvl=index)
            _element(lvl, "w:start", val=start if index == level else 1)
            _element(lvl, "w:numFmt", val=LIST_FORMATS[kind])
            bullet = {"disc": "•", "circle": "◦", "square": "▪", "none": ""}.get(kind, f"%{index + 1}.")
            _element(lvl, "w:lvlText", val=bullet)
            _element(lvl, "w:lvlJc", val="left")
            ppr = _element(lvl, "w:pPr")
            _element(ppr, "w:ind", left=360 * (index + 1), hanging=180)
        num = _element(root, "w:num", numId=self.next_num)
        _element(num, "w:abstractNumId", val=self.next_abstract)
        override = _element(num, "w:lvlOverride", ilvl=level)
        _element(override, "w:startOverride", val=start)
        self.next_num += 1
        self.next_abstract += 1
        return self.next_num - 1

    def list(self, node, flow):
        level = sum(parent.tag in {"ul", "ol"} for parent in node.iterancestors())
        style = self.source.styles[node]
        kind = style.get("list-style-type")
        if kind not in LIST_FORMATS:
            kind = "decimal" if node.tag == "ol" else "disc"
        number = self.numbering(kind, level, int(node.get("start", 1)))
        paragraphs = []
        for item in node:
            if self.source.hidden(item):
                continue
            if item.get("value") is not None and node.tag == "ol":
                number = self.numbering(kind, level, int(item.get("value")))
            item_flow = _Flow(self, flow.container, flow.width, self.source.styles[item], flow.indent,
                              flow.right, (number, level))
            item_flow.pending_before, flow.pending_before = flow.pending_before, 0
            self.children(item, item_flow)
            if not item_flow.paragraphs:
                item_flow.ensure(self.source.styles[item])
            if item_flow.paragraphs:
                self.after(self.source.styles[item], item_flow.paragraphs[-1])
            flow.paragraphs.extend(item_flow.paragraphs)
            paragraphs.extend(item_flow.paragraphs)
        if paragraphs:
            first, last = paragraphs[0], paragraphs[-1]
            top = _length(style.get("margin-top", "0"), reference=flow.width) or 0
            bottom = _length(style.get("margin-bottom", "0"), reference=flow.width) or 0
            first.paragraph_format.space_before = Pt(max(first.paragraph_format.space_before.pt if first.paragraph_format.space_before else 0, min(PAGE_HEIGHT, top)))
            last.paragraph_format.space_after = Pt(max(last.paragraph_format.space_after.pt if last.paragraph_format.space_after else 0, min(PAGE_HEIGHT, bottom)))
            if style.get("break-before", style.get("page-break-before")) in {"page", "always", "left", "right", "recto", "verso"}:
                first.paragraph_format.page_break_before = True
            self.after(style, last)

    def borders(self, style):
        if not any(key.startswith("border-") and key.endswith(("-style", "-width", "-color")) for key in style):
            return None
        borders = OxmlElement("w:tcBorders")
        for side in ("top", "left", "bottom", "right"):
            kind = style.get(f"border-{side}-style", "none")
            value = {"solid": "single", "double": "double", "dotted": "dotted", "dashed": "dashed",
                     "groove": "single", "ridge": "single", "inset": "single", "outset": "single"}.get(kind, "nil")
            width = style.get(f"border-{side}-width", "1px")
            size = {"thin": .75, "medium": 2.25, "thick": 3.75}.get(width, _length(width))
            if size == 0:
                value = "nil"
            size = size or .75
            _element(borders, f"w:{side}", val=value, sz=max(2, min(96, round(size * 8))),
                     color=_color(style.get(f"border-{side}-color", "black")) or "000000")
        return borders

    def table(self, node, flow):
        rows, count, cells = self.source.tables[node]
        style = self.source.styles[node]
        if node.get("id"):
            self.bookmark(node, flow)
            flow.ensure().paragraph_format.keep_with_next = True
            flow.end()
        for child in node:
            if child.tag == "caption":
                paragraph = flow.ensure(self.source.styles[child])
                self.children(child, flow)
                paragraph.paragraph_format.keep_with_next = True
                flow.end()
        before = style.get("break-before", style.get("page-break-before"))
        if before in {"always", "page", "left", "right", "recto", "verso"}:
            flow.ensure().paragraph_format.page_break_before = True
            flow.end()
        available = max(12, flow.width - flow.indent - flow.right)
        requested = _length(style.get("width", node.get("width", "")), reference=available)
        width = min(available, requested) if requested and requested > 0 else available
        maximum = _length(style.get("max-width", ""), reference=available)
        if maximum and maximum > 0:
            width = min(width, maximum)
        widths = [None] * count
        column_index = 0
        for child in node:
            if child.tag not in {"colgroup", "col"}:
                continue
            columns = list(child) if child.tag == "colgroup" and len(child) else [child]
            for column in columns:
                try:
                    span = int(column.get("span", 1))
                    if span < 1 or column_index + span > count:
                        raise ValueError()
                except ValueError:
                    self.source.note("errors", "Invalid table column span.", column)
                    return
                value = self.source.styles[column].get("width", column.get("width", self.source.styles[child].get("width", "")))
                amount = _length(value, reference=width)
                for _ in range(span):
                    widths[column_index] = amount if amount and amount > 0 else None
                    column_index += 1
        # Explicit single-column widths take precedence over spanning headers.
        for cell, row, col, rowspan, colspan in sorted(cells, key=lambda entry: entry[4]):
            amount = _length(self.source.styles[cell].get("width", cell.get("width", "")), reference=width)
            if amount and amount > 0:
                unknown = [index for index in range(col, col + colspan) if widths[index] is None]
                remaining = amount - sum(widths[index] or 0 for index in range(col, col + colspan))
                for index in range(col, col + colspan):
                    if widths[index] is None:
                        widths[index] = remaining / len(unknown) if remaining > 0 else amount / colspan
        known = sum(value or 0 for value in widths)
        missing = sum(value is None for value in widths)
        fallback = max(width / count / 4, (width - known) / missing) if missing else width / count
        widths = [value if value is not None else fallback for value in widths]
        factor = width / sum(widths)
        widths = [max(.1, value * factor) for value in widths]
        table = flow.container.add_table(rows=len(rows), cols=count)
        self.source.report["tableCount"] = len(self.source.tables)
        table.autofit = False
        table.alignment = {"center": WD_TABLE_ALIGNMENT.CENTER, "right": WD_TABLE_ALIGNMENT.RIGHT}.get(style.get("text-align"), WD_TABLE_ALIGNMENT.LEFT)
        if style.get("margin-left") == style.get("margin-right") == "auto":
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
        props = table._tbl.tblPr
        table_width = props.find(qn("w:tblW"))
        table_width.set(qn("w:type"), "dxa")
        table_width.set(qn("w:w"), str(round(width * 20)))
        if flow.indent:
            indentation = _element(props, "w:tblInd", w=round(flow.indent * 20), type="dxa")
            props.insert_element_before(indentation, "w:tblBorders", "w:shd", "w:tblLayout", "w:tblCellMar", "w:tblLook")
        table_borders = self.borders(style)
        if table_borders is None and node.get("border") not in {None, "0"}:
            table_borders = self.borders({f"border-{side}-style": "solid" for side in ("top", "bottom", "left", "right")})
        if table_borders is not None:
            table_borders.tag = qn("w:tblBorders")
            # The HTML border attribute also draws internal cell boundaries.
            if node.get("border") not in {None, "0"}:
                for axis in ("insideH", "insideV"):
                    _element(table_borders, "w:" + axis, val="single", sz=4, color="000000")
            props.insert_element_before(table_borders, "w:shd", "w:tblLayout", "w:tblCellMar", "w:tblLook")
        for index, column in enumerate(table.columns):
            column.width = Pt(widths[index])
        grid = [[_Cell(tc, table) for tc in tr.tc_lst] for tr in table._tbl.tr_lst]
        for row in grid:
            for index, cell in enumerate(row):
                cell.width = Pt(widths[index])
        header = True
        for index, original_row in enumerate(rows):
            is_header = original_row.getparent().tag == "thead" or all(child.tag == "th" for child in original_row)
            row_props = table.rows[index]._tr.get_or_add_trPr()
            if header and is_header:
                _element(row_props, "w:tblHeader", val="true")
            else:
                header = False
            row_style = self.source.styles[original_row]
            if row_style.get("break-inside", row_style.get("page-break-inside")) in {"avoid", "avoid-page"}:
                cannot_split = _element(row_props, "w:cantSplit")
                row_props.insert_element_before(cannot_split, "w:trHeight", "w:tblHeader")
        for original, row, col, rowspan, colspan in cells:
            cell = grid[row][col]
            if rowspan > 1 or colspan > 1:
                cell = cell.merge(grid[row + rowspan - 1][col + colspan - 1])
            cell_width = sum(widths[col:col + colspan])
            cell.width = Pt(cell_width)
            cell_style = self.source.styles[original]
            properties = cell._tc.get_or_add_tcPr()
            row_style = self.source.styles[rows[row]]
            group_style = self.source.styles.get(rows[row].getparent(), {})
            background = cell_style.get("background-color", row_style.get("background-color", group_style.get("background-color", style.get("background-color", ""))))
            fill = _color(background)
            if fill:
                _element(properties, "w:shd", fill=fill, val="clear")
            borders = self.borders(cell_style)
            if borders is not None:
                properties.insert_element_before(borders, "w:shd", "w:tcMar", "w:vAlign")
            cell.vertical_alignment = {"middle": WD_CELL_VERTICAL_ALIGNMENT.CENTER, "bottom": WD_CELL_VERTICAL_ALIGNMENT.BOTTOM}.get(cell_style.get("vertical-align", row_style.get("vertical-align")), WD_CELL_VERTICAL_ALIGNMENT.TOP)
            padding = []
            margins = OxmlElement("w:tcMar")
            for side in ("top", "left", "bottom", "right"):
                amount = _length(cell_style.get(f"padding-{side}", ""), reference=cell_width)
                if amount is not None:
                    amount = max(0, min(cell_width / 4, amount))
                    _element(margins, f"w:{side}", w=round(amount * 20), type="dxa")
                if side in {"left", "right"}:
                    padding.append(amount if amount is not None else 5.4)
            if len(margins):
                properties.insert_element_before(margins, "w:textDirection", "w:tcFitText", "w:vAlign")
            # Cell box formatting belongs on tcPr; explicit child formatting
            # still applies to that child's own paragraph or text run.
            paragraph_style = {key: value for key, value in cell_style.items()
                               if not key.startswith(("border-", "padding-")) and key != "background-color"}
            cell_flow = _Flow(self, cell, max(1, cell_width - sum(padding)), paragraph_style)
            self.children(original, cell_flow)
            if not cell_flow.paragraphs:
                cell_flow.ensure()
        top = _length(style.get("margin-top", "0"), reference=width) or 0
        bottom = _length(style.get("margin-bottom", "0"), reference=width) or 0
        if top > 0 and flow.paragraphs:
            previous = flow.paragraphs[-1].paragraph_format
            previous.space_after = Pt(max(previous.space_after.pt if previous.space_after else 0, min(PAGE_HEIGHT, top)))
        flow.pending_before = max(flow.pending_before, min(PAGE_HEIGHT, bottom))
        if style.get("break-after", style.get("page-break-after")) in {"always", "page", "left", "right", "recto", "verso"}:
            flow.ensure().add_run().add_break(WD_BREAK.PAGE)
            flow.end()


def _convert(source, resources, config, export):
    parsed = None
    try:
        parsed = _Source(source, resources, config)
        if parsed.report["errors"]:
            return b"", parsed.report
        document = _Renderer(parsed).render()
        if parsed.report["errors"] or not export:
            return b"", parsed.report
        output = BytesIO()
        document.save(output)
        return output.getvalue(), parsed.report
    except (ValueError, TypeError, KeyError, OverflowError, RecursionError, etree.LxmlError, OSError) as exc:
        report = parsed.report if parsed is not None else {"title": "", "paragraphCount": 0, "tableCount": 0,
                 "imageCount": 0, "warnings": [], "errors": [], "resources": []}
        report["errors"].append(f"HTML cannot be converted reliably: {type(exc).__name__}: {exc}")
        return b"", report


def inspect_html(source: bytes, resources: dict[str, bytes] | None = None, config: dict | None = None) -> dict:
    """Discover raw local image src values; a supplied dict also validates images."""
    return _convert(source, resources, config, False)[1]


def export_docx(source: bytes, resources: dict[str, bytes] | None = None, config: dict | None = None) -> tuple[bytes, dict]:
    """Return editable DOCX bytes and report, or empty bytes if validation fails."""
    return _convert(source, resources if resources is not None else {}, config, True)
