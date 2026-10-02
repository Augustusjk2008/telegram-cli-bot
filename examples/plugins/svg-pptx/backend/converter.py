"""One SVG viewport -> one slide, interleaving raster graphics and native text."""

from __future__ import annotations

import copy
import importlib.util
import math
from io import BytesIO


DEPENDENCIES = {"pptx": "python-pptx", "resvg_py": "resvg_py", "PIL": "Pillow",
                "tinycss2": "tinycss2", "cssselect2": "cssselect2", "defusedxml": "defusedxml"}


def dependency_status() -> list[str]:
    return [package for module, package in DEPENDENCIES.items() if importlib.util.find_spec(module) is None]


def _scene(source: bytes, config: dict | None):
    missing = dependency_status()
    if missing:
        raise ValueError("缺少依赖: " + ", ".join(missing) + "；请使用宿主 Python 安装插件 requirements.txt")
    from svg_parser import parse
    return parse(source, config)


def inspect_svg(source: bytes, config: dict | None = None) -> dict:
    try:
        return _scene(source, config).report()
    except Exception as exc:
        return {"width": 0, "height": 0, "textCount": 0, "layerCount": 0,
                "warnings": [], "errors": [str(exc)]}


def graphics_svg(scene, layer) -> str:
    """Freeze styles before removing text so structural CSS selectors stay stable."""
    import xml.etree.ElementTree as ET
    import tinycss2
    from svg_parser import DEFINITIONS, GRAPHICS, SVG, tag

    root = copy.deepcopy(scene.root)
    originals, copies = list(scene.root.iter()), list(root.iter())
    parents = {child: parent for parent in originals for child in parent}
    copy_parents = {child: parent for parent in copies for child in parent}
    selected = set(layer.elements)
    ids = {element.get("id") for element in originals if element.get("id")}

    def in_definition(element):
        while element is not None:
            if tag(element) in DEFINITIONS:
                return True
            element = parents.get(element)
        return False

    # A painted element can be reused later by <use>. Give the resource its own
    # definition, so hiding a different paint layer does not hide that instance.
    referenced = {target for targets in scene.references.values() for target in targets if not in_definition(target)}
    replacements = {}
    resources = ET.Element(f"{{{SVG}}}defs")
    for index, target in enumerate(element for element in originals if element in referenced):
        new_id = f"svg_pptx_resource_{index}"
        while new_id in ids:
            new_id += "_"
        ids.add(new_id)
        replacements[target.get("id")] = new_id
        resource = copy.deepcopy(target)
        for ordinal, (original, child) in enumerate(zip(target.iter(), resource.iter())):
            child.set("style", ";".join(f"{key}:{value}" for key, value in scene.specified_styles[original].items()))
            if child.get("id"):
                child_id = new_id if ordinal == 0 else f"{new_id}_child_{ordinal}"
                if ordinal:
                    while child_id in ids:
                        child_id += "_"
                    ids.add(child_id)
                child.set("id", child_id)
        resources.append(resource)
    for original, copied in zip(originals, copies):
        style = scene.specified_styles[original] if in_definition(original) else scene.styles[original]
        copied.set("style", ";".join(f"{key}:{value}" for key, value in style.items()))
        if tag(original) in {"text", "tspan", "textPath", "style"}:
            parent = copy_parents.get(copied)
            if parent is not None and copied in list(parent):
                parent.remove(copied)
        elif tag(original) in GRAPHICS and original not in selected:
            ancestor = parents.get(original)
            while ancestor is not None and tag(ancestor) not in DEFINITIONS:
                ancestor = parents.get(ancestor)
            if ancestor is None:
                copied.set("style", copied.get("style", "") + ";display:none!important")
    root.append(resources)
    for parent in list(root.iter()):
        for child in list(parent):
            if tag(child) in {"text", "tspan", "textPath", "style"}:
                parent.remove(child)

    def rewrite(tokens):
        for token in tokens:
            if token.type == "url" and token.value.startswith("#") and token.value[1:] in replacements:
                token.representation = f'url("#{replacements[token.value[1:]]}")'
            elif token.type == "function":
                if token.lower_name == "url":
                    for argument in token.arguments:
                        if argument.type == "string" and argument.value.startswith("#") and argument.value[1:] in replacements:
                            argument.representation = f'"#{replacements[argument.value[1:]]}"'
                rewrite(token.arguments)

    for element in root.iter():
        for key, value in list(element.attrib.items()):
            if key.split("}")[-1] == "href" and value.startswith("#") and value[1:] in replacements:
                element.set(key, "#" + replacements[value[1:]])
            elif "url" in value.lower() or "\\" in value:
                tokens = tinycss2.parse_component_value_list(value)
                rewrite(tokens)
                element.set(key, tinycss2.serialize(tokens))
    root.set("width", f"{scene.width:g}px")
    root.set("height", f"{scene.height:g}px")
    ET.register_namespace("", "http://www.w3.org/2000/svg")
    return ET.tostring(root, encoding="unicode")


def _textbox(slide, box, units: float, number: int) -> None:
    from pptx.dml.color import RGBColor
    from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
    from pptx.oxml.xmlchemy import OxmlElement
    from pptx.oxml.ns import qn
    from pptx.util import Emu, Pt

    scale, angle = box.matrix.similarity()
    ascent = max(run.ascent for run in box.runs)
    descent = max(run.descent for run in box.runs)
    offset = {"auto": 0, "alphabetic": 0,
              "middle": max(run.x_height for run in box.runs) / 2,
              "central": (ascent - descent) / 2,
              "text-before-edge": ascent, "text-after-edge": -descent}[box.baseline]
    local_height, local_width = ascent + descent, max(box.width, 0.1)
    center_x, center_y = box.matrix.point(box.x + local_width / 2, box.y + offset - ascent + local_height / 2)
    width, height = local_width * scale * units, local_height * scale * units
    shape = slide.shapes.add_textbox(round(center_x * units - width / 2), round(center_y * units - height / 2),
                                    max(1, round(width)), max(1, round(height)))
    shape.name = f"SVG text {number}"
    shape.rotation = angle % 360
    frame = shape.text_frame
    frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
    frame.word_wrap = False
    frame.auto_size = MSO_AUTO_SIZE.NONE
    frame.vertical_anchor = MSO_ANCHOR.TOP
    paragraph = frame.paragraphs[0]
    paragraph.alignment = {"start": PP_ALIGN.LEFT, "middle": PP_ALIGN.CENTER, "end": PP_ALIGN.RIGHT}[box.anchor]
    paragraph.space_before = paragraph.space_after = Pt(0)
    paragraph.line_spacing = Emu(round(height))
    paragraph.font.size = round(max(run.size for run in box.runs) * scale * units)
    paragraph.font.name = box.runs[0].face.family
    paragraph._p.get_or_add_pPr().append(OxmlElement("a:buNone"))
    for source in box.runs:
        run = paragraph.add_run()
        run.text = source.text
        run.font.size = max(1, round(source.size * scale * units))
        run.font.name = source.face.family
        run.font.bold, run.font.italic = source.bold, source.italic
        run.font.underline = "underline" in source.decoration
        run.font.color.rgb = RGBColor(*source.color)
        properties = run._r.get_or_add_rPr()
        properties.set("lang", "en-US")
        properties.set("altLang", "zh-CN")
        properties.set("strike", "sngStrike" if "line-through" in source.decoration else "noStrike")
        for name in ("a:ea", "a:cs"):
            child = OxmlElement(name)
            child.set("typeface", source.face.family)
            properties.append(child)
        if source.alpha < 1:
            alpha = OxmlElement("a:alpha")
            alpha.set("val", str(round(source.alpha * 100_000)))
            properties.find(qn("a:solidFill")).find(qn("a:srgbClr")).append(alpha)


def export_pptx(source: bytes, config: dict | None = None) -> tuple[bytes, dict]:
    import resvg_py
    from pptx import Presentation
    from svg_parser import Layer, limit

    config = config or {}
    scene = _scene(source, config)
    if scene.errors:
        raise ValueError("\n".join(scene.errors))
    presentation = Presentation()
    # PowerPoint accepts 1..56 inches. Scale only when required, preserving aspect.
    units = 9525 * max(96 / min(scene.width, scene.height), min(1, 5376 / max(scene.width, scene.height)))
    presentation.slide_width, presentation.slide_height = round(scene.width * units), round(scene.height * units)
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    render_side = min(limit(config, "renderMaxSide", 4096, 4096), max(1, math.ceil(max(scene.width, scene.height) * 3)))
    zoom = render_side / max(scene.width, scene.height)
    render_width, render_height = max(1, round(scene.width * zoom)), max(1, round(scene.height * zoom))
    image_bytes, text_number = 0, 0
    for item in scene.items:
        if isinstance(item, Layer):
            picture = resvg_py.svg_to_bytes(svg_string=graphics_svg(scene, item), width=render_width, height=render_height,
                                           skip_system_fonts=True, log_information=False)
            image_bytes += len(picture)
            if image_bytes > 32 * 1024 * 1024:
                raise ValueError("图片层总大小超过 32 MiB 产物限制")
            slide.shapes.add_picture(BytesIO(picture), 0, 0, presentation.slide_width, presentation.slide_height)
        else:
            text_number += 1
            _textbox(slide, item, units, text_number)
    output = BytesIO()
    presentation.save(output)
    content = output.getvalue()
    if len(content) > 32 * 1024 * 1024:
        raise ValueError("PPTX 超过 32 MiB 产物限制")
    return content, scene.report()
